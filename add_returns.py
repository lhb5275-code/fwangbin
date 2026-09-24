#!/usr/bin/env python3
"""sp500_last5q 결과 Excel에 최근 1년 주가수익률과 수익률 분석 시트를 추가한다.

사용법:
    python add_returns.py 입력.xlsx [--asof-file sp500_asof20250924_last5q_*.xlsx] [--out 출력.xlsx]

추가 내용 ('실적' 시트 오른쪽 열):
    GICS 섹터 | 시작일 종가 | 종료일 종가 | 1년 주가수익률(수식) | 1년 총수익률(배당 포함)
    | 1년 전 매출성장률(YoY) | 1년 전 선정 기준          (← --asof-file 을 주면)
새 시트 '수익률 분석': 전체 요약, 섹터별, 수익률 5분위별 특징, 상위 30 종목.

가격: Yahoo Finance(yfinance) 일별 종가. 주가수익률은 분할 조정 종가(Close), 총수익률은 배당까지 조정한
Adj Close 기준. 종료일은 미국 장 마감 전이면 전 거래일 종가를 쓴다.
"""

import argparse
import os
import sys
from datetime import datetime, timedelta
from io import StringIO
from zoneinfo import ZoneInfo

import numpy as np
import pandas as pd
import requests
import yfinance as yf
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

WIKI_URL = "https://en.wikipedia.org/wiki/List_of_S%26P_500_companies"
GITHUB_SP500_CSV = "https://raw.githubusercontent.com/datasets/s-and-p-500-companies/main/data/constituents.csv"
ET = ZoneInfo("America/New_York")
FMT_PX = "#,##0.00"
FMT_PCT = "0.0%"
HDR_FILL = PatternFill("solid", fgColor="DDEBF7")
BOLD = Font(bold=True)


def gics_sectors():
    try:
        r = requests.get(WIKI_URL, headers={"User-Agent": "Mozilla/5.0 (sp500-last5q script)"}, timeout=30)
        r.raise_for_status()
        df = pd.read_html(StringIO(r.text), attrs={"id": "constituents"})[0]
    except Exception:  # noqa: BLE001
        df = pd.read_csv(GITHUB_SP500_CSV)
    t = df["Symbol"].astype(str).str.replace(".", "-", regex=False)
    return dict(zip(t, df["GICS Sector"]))


def price_window(tickers, start, end_asof):
    """시작일(start) 종가와 종료일(마지막 완결 거래일) 종가를 Close / Adj Close 로 반환."""
    px = yf.download(
        tickers + ["^GSPC"],
        start=(pd.Timestamp(start) - timedelta(days=7)).strftime("%Y-%m-%d"),
        end=(pd.Timestamp(end_asof) + timedelta(days=1)).strftime("%Y-%m-%d"),
        auto_adjust=False,
        group_by="column",
        progress=False,
        threads=True,
    )
    close, adj = px["Close"], px["Adj Close"]
    now = datetime.now(ET)
    last = close.index.max()
    # 미국 장 마감(16:00 ET) 전이면 당일 행은 장중 가격이므로 제외
    if last.date() == now.date() and now.hour < 16:
        close, adj = close.loc[close.index < last], adj.loc[adj.index < last]
    end_day = close.index.max()
    start_day = pd.Timestamp(start)
    if start_day not in close.index:
        raise SystemExit(f"{start} 은(는) 거래일이 아닙니다")
    return close, adj, start_day, end_day


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx", help="sp500_last5q.py 결과 파일")
    ap.add_argument("--asof-file", help="--as-of 로 만든 과거 시점 파일 (1년 전 성장률·선정 기준 비교용)")
    ap.add_argument("--start", default=None, help="수익률 시작일 (기본: 오늘로부터 1년 전 같은 날짜)")
    ap.add_argument("--out", help="출력 파일 (기본: 입력파일명_returns.xlsx)")
    args = ap.parse_args()

    today = datetime.now(ET).date()
    start = args.start or today.replace(year=today.year - 1).strftime("%Y-%m-%d")
    out = args.out or os.path.splitext(args.xlsx)[0] + "_returns.xlsx"

    wb = load_workbook(args.xlsx)
    ws = wb["실적"]
    hdr = [c.value for c in ws[1]]
    if "1년 주가수익률" in " ".join(str(h) for h in hdr):
        raise SystemExit("이미 수익률 열이 있는 파일입니다")
    col = {h: i + 1 for i, h in enumerate(hdr)}
    n_rows = ws.max_row
    tickers = [ws.cell(r, col["티커"]).value for r in range(2, n_rows + 1)]

    print(f"가격 수집: {len(tickers)}개 종목, {start} ~ 최근 거래일")
    close, adj, start_day, end_day = price_window(tickers, start, today.strftime("%Y-%m-%d"))
    s_lbl, e_lbl = start_day.strftime("%Y-%m-%d"), end_day.strftime("%Y-%m-%d")
    sectors = gics_sectors()

    asof = {}
    asof_lbl = None
    if args.asof_file:
        aw = load_workbook(args.asof_file, data_only=True)["실적"]
        ah = {c.value: i for i, c in enumerate(aw[1])}
        for r in aw.iter_rows(min_row=2, values_only=True):
            asof[r[ah["티커"]]] = (r[ah["매출성장률(YoY)"]], r[ah["선정 기준"]])
        base = os.path.basename(args.asof_file)
        asof_lbl = base.split("asof")[1][:8] if "asof" in base else "과거"
        asof_lbl = f"{asof_lbl[:4]}-{asof_lbl[4:6]}-{asof_lbl[6:8]}" if asof_lbl.isdigit() else asof_lbl

    # ---- 실적 시트에 열 추가
    c0 = ws.max_column + 1
    new_hdr = ["GICS 섹터", f"주가 {s_lbl} ($)", f"주가 {e_lbl} ($)", "1년 주가수익률", "1년 총수익률(배당 포함)"]
    if asof:
        new_hdr += [
            f"1년 전 매출성장률(YoY, {asof_lbl} 시점)",
            f"1년 전 선정 기준 ({asof_lbl} 시점)",
            "매출 YoY 변화(%p, 현재-1년 전)",
        ]
    C = {h: c0 + i for i, h in enumerate(new_hdr)}
    L = {h: get_column_letter(c) for h, c in C.items()}
    for h, c in C.items():
        cell = ws.cell(1, c, h)
        cell.font, cell.fill = BOLD, HDR_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    no_start = []
    for r, t in enumerate(tickers, start=2):
        fill = ws.cell(r, 1).fill
        p0 = close.at[start_day, t] if t in close else np.nan
        p1 = close.at[end_day, t] if t in close else np.nan
        a0 = adj.at[start_day, t] if t in adj else np.nan
        a1 = adj.at[end_day, t] if t in adj else np.nan
        if pd.isna(p0):
            no_start.append(t)
        vals = {
            "GICS 섹터": sectors.get(t),
            f"주가 {s_lbl} ($)": None if pd.isna(p0) else round(float(p0), 4),
            f"주가 {e_lbl} ($)": None if pd.isna(p1) else round(float(p1), 4),
            "1년 주가수익률": f'=IF(AND(ISNUMBER({L[new_hdr[1]]}{r}),ISNUMBER({L[new_hdr[2]]}{r})),'
            f'{L[new_hdr[2]]}{r}/{L[new_hdr[1]]}{r}-1,"")',
            "1년 총수익률(배당 포함)": None if pd.isna(a0) or pd.isna(a1) else float(a1 / a0 - 1),
        }
        if asof:
            yoy, tier = asof.get(t, (None, None))
            vals[new_hdr[5]] = yoy
            vals[new_hdr[6]] = tier
            cur, prev = f"{get_column_letter(col['매출성장률(YoY)'])}{r}", f"{L[new_hdr[5]]}{r}"
            vals[new_hdr[7]] = f'=IF(AND(ISNUMBER({cur}),ISNUMBER({prev})),{cur}-{prev},"")' 
        for h, v in vals.items():
            cell = ws.cell(r, C[h], v)
            if fill is not None and fill.fill_type:
                cell.fill = PatternFill("solid", fgColor=fill.fgColor.rgb)
            if h.startswith("주가"):
                cell.number_format = FMT_PX
            elif "수익률" in h or "성장률" in h or "변화" in h:
                cell.number_format = FMT_PCT
            if isinstance(v, str) and h.startswith("1년 전 매출"):
                cell.alignment = Alignment(horizontal="right")
    last_col = get_column_letter(ws.max_column)
    ws.auto_filter.ref = f"A1:{last_col}{n_rows}"
    for h, c in C.items():
        ws.column_dimensions[get_column_letter(c)].width = max(12, min(40, len(h) * 1.6))

    # ---- 수익률 분석 시트
    if "수익률 분석" in wb.sheetnames:
        del wb["수익률 분석"]
    an = wb.create_sheet("수익률 분석", 1)
    rng = lambda h: f"'실적'!${L[h]}$2:${L[h]}${n_rows}"  # noqa: E731
    ret, sec = rng("1년 주가수익률"), rng("GICS 섹터")
    tier = f"'실적'!${get_column_letter(col['선정 기준'])}$2:${get_column_letter(col['선정 기준'])}${n_rows}"
    ryoy = f"'실적'!${get_column_letter(col['매출성장률(YoY)'])}$2:${get_column_letter(col['매출성장률(YoY)'])}${n_rows}"
    oyoy = f"'실적'!${get_column_letter(col['영업이익성장률(YoY)'])}$2:${get_column_letter(col['영업이익성장률(YoY)'])}${n_rows}"
    qoq0 = f"'실적'!${get_column_letter(col['매출 QoQ(Q0)'])}$2:${get_column_letter(col['매출 QoQ(Q0)'])}${n_rows}"
    pyoy = rng(new_hdr[5]) if asof else None
    ptier = rng(new_hdr[6]) if asof else None
    dyoy = rng(new_hdr[7]) if asof else None

    def title(row, text):
        an.cell(row, 1, text).font = Font(bold=True, size=12)

    def header(row, labels):
        for i, h in enumerate(labels, start=1):
            c = an.cell(row, i, h)
            c.font, c.fill = BOLD, HDR_FILL
            c.alignment = Alignment(horizontal="center", wrap_text=True)

    def med(cond, vals):
        return f"=IFERROR(MEDIAN(IF(({cond})*ISNUMBER({vals}),{vals})),\"\")"

    r = 1
    title(r, f"S&P 500 최근 1년 주가수익률 분석 ({s_lbl} 종가 → {e_lbl} 종가)")
    r += 2
    header(r, ["항목", "값", "설명"])
    spx0, spx1 = close.at[start_day, "^GSPC"], close.at[end_day, "^GSPC"]
    summary = [
        ("수익률 계산 종목 수", f"=COUNT({ret})", "시작일 이전 상장 종목 (이후 상장·분사 종목 제외)", "0"),
        ("S&P 500 지수 수익률 (^GSPC)", float(spx1 / spx0 - 1), f"지수 종가 {spx0:,.2f} → {spx1:,.2f} (Yahoo)", FMT_PCT),
        ("종목 수익률 평균 (동일가중)", f"=AVERAGE({ret})", "", FMT_PCT),
        ("종목 수익률 중앙값", f"=MEDIAN({ret})", "", FMT_PCT),
        ("상승 종목 비율", f'=COUNTIF({ret},">0")/COUNT({ret})', "", FMT_PCT),
        ("+50% 이상 종목 수", f'=COUNTIF({ret},">=0.5")', "", "0"),
        ("+100% 이상 종목 수", f'=COUNTIF({ret},">=1")', "", "0"),
    ]
    for lbl, v, note, fmt in summary:
        r += 1
        an.cell(r, 1, lbl)
        c = an.cell(r, 2, v)
        c.number_format = fmt
        an.cell(r, 3, note)

    # 섹터별
    r += 3
    title(r, "섹터별 수익률")
    r += 1
    header(r, ["GICS 섹터", "종목 수", "평균 수익률", "중앙값 수익률", "+50% 이상 종목 수", "+50% 이상 비율", "중앙값 매출 YoY(현재)"])
    for s in sorted({v for v in sectors.values() if isinstance(v, str)}):
        r += 1
        an.cell(r, 1, s)
        an.cell(r, 2, f"=COUNTIFS({sec},A{r},{ret},\">-100\")")
        an.cell(r, 3, f"=IFERROR(AVERAGEIFS({ret},{sec},A{r}),\"\")").number_format = FMT_PCT
        an.cell(r, 4).value = ArrayFormula(f"D{r}", med(f"{sec}=A{r}", ret))
        an.cell(r, 4).number_format = FMT_PCT
        an.cell(r, 5, f"=COUNTIFS({sec},A{r},{ret},\">=0.5\")")
        an.cell(r, 6, f"=IFERROR(E{r}/B{r},\"\")").number_format = FMT_PCT
        an.cell(r, 7).value = ArrayFormula(f"G{r}", med(f"{sec}=A{r}", ryoy))
        an.cell(r, 7).number_format = FMT_PCT
    an.cell(r + 1, 1, "※ 섹터 행은 필터로 정렬 가능. 평균은 극단값 영향이 크므로 중앙값을 함께 볼 것").font = Font(italic=True)

    # 5분위
    r += 3
    title(r, "수익률 5분위별 특징 (1분위 = 수익률 상위 20%)")
    r += 1
    labels = ["분위", "수익률 하한", "수익률 상한", "종목 수", "중앙값 수익률", "중앙값 매출 YoY(현재)", "중앙값 영업이익 YoY(현재)",
              "중앙값 최근 매출 QoQ(현재)", "현재 선정(①②) 비율"]
    if asof:
        labels += [f"중앙값 매출 YoY({asof_lbl})", f"{asof_lbl} 선정(①②) 비율", "중앙값 매출 YoY 변화(%p)"]
    header(r, labels)
    for k in range(5):
        r += 1
        hi_p, lo_p = 1 - 0.2 * k, 0.8 - 0.2 * k
        an.cell(r, 1, f"{k + 1}분위")
        an.cell(r, 2, f"=PERCENTILE({ret},{lo_p:.1f})").number_format = FMT_PCT
        an.cell(r, 3, f"=PERCENTILE({ret},{hi_p:.1f})").number_format = FMT_PCT
        upper = f"{ret}<=C{r}" if k == 0 else f"{ret}<C{r}"
        cond = f"ISNUMBER({ret})*({ret}>=B{r})*({upper})"
        crit_hi = f'"<="&C{r}' if k == 0 else f'"<"&C{r}'
        an.cell(r, 4, f'=COUNTIFS({ret},">="&B{r},{ret},{crit_hi})')
        cells = [(5, ret), (6, ryoy), (7, oyoy), (8, qoq0)]
        if asof:
            cells += [(10, pyoy), (12, dyoy)]
        for cc, vals in cells:
            ref = f"{get_column_letter(cc)}{r}"
            an.cell(r, cc).value = ArrayFormula(ref, med(cond, vals))
            an.cell(r, cc).number_format = FMT_PCT
        an.cell(r, 9).value = ArrayFormula(f"I{r}", f'=IFERROR(SUM(({cond})*({tier}<>""))/D{r},"")')
        an.cell(r, 9).number_format = FMT_PCT
        if asof:
            an.cell(r, 11).value = ArrayFormula(f"K{r}", f'=IFERROR(SUM(({cond})*({ptier}<>""))/D{r},"")')
            an.cell(r, 11).number_format = FMT_PCT
    an.cell(r + 1, 1, "※ 영업이익 YoY 중앙값은 숫자 값만 사용 ('흑자전환'·'N/M' 등 제외). "
            "'현재' = 이 파일의 최근 5개 분기, 1년 전 = SEC 제출 자료 기준 당시 최근 5개 분기").font = Font(italic=True)

    # 상위 30
    r += 4
    title(r, "1년 주가수익률 상위 30 종목")
    r += 1
    top_hdr = ["순위", "티커", "기업명", "GICS 섹터", "1년 주가수익률", "매출 YoY(현재)", "최근 QoQ(현재)", "현재 선정 기준"]
    if asof:
        top_hdr += [f"매출 YoY({asof_lbl})", f"{asof_lbl} 선정 기준"]
    top_hdr += ["주요 사업"]
    header(r, top_hdr)
    rets = []
    for i, t in enumerate(tickers, start=2):
        p0 = close.at[start_day, t] if t in close else np.nan
        p1 = close.at[end_day, t] if t in close else np.nan
        if pd.notna(p0) and pd.notna(p1):
            rets.append((p1 / p0 - 1, i))
    rets.sort(reverse=True)
    ko = {}
    ko_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "business_ko.csv")
    if os.path.exists(ko_path):
        kdf = pd.read_csv(ko_path, dtype=str, encoding="utf-8-sig").dropna()
        ko = dict(zip(kdf["ticker"], kdf["business_ko"]))
    for rank, (_, src) in enumerate(rets[:30], start=1):
        r += 1
        t = ws.cell(src, col["티커"]).value
        tcol = get_column_letter(col["티커"])
        refs = [col["기업명"], C["GICS 섹터"], C["1년 주가수익률"], col["매출성장률(YoY)"], col["매출 QoQ(Q0)"], col["선정 기준"]]
        if asof:
            refs += [C[new_hdr[5]], C[new_hdr[6]]]
        refs += [col["주요 사업"]]
        an.cell(r, 1, rank)
        an.cell(r, 2, t)
        for j, cref in enumerate(refs, start=3):
            lk = f"INDEX('실적'!${get_column_letter(cref)}$2:${get_column_letter(cref)}${n_rows},MATCH($B{r},'실적'!${tcol}$2:${tcol}${n_rows},0))"
            if cref == col["주요 사업"] and not ws.cell(src, cref).value and t in ko:
                c = an.cell(r, j, ko[t])  # 실적 시트에 설명이 없는 종목은 business_ko.csv 값
            else:
                c = an.cell(r, j, f'=IF({lk}="","",{lk})')
            if top_hdr[j - 1] in ("1년 주가수익률",) or "YoY" in top_hdr[j - 1] or "QoQ" in top_hdr[j - 1]:
                c.number_format = FMT_PCT
    for j, w in enumerate([22, 12, 22, 16, 14, 14, 14, 22, 14, 22, 50, 14], start=1):
        an.column_dimensions[get_column_letter(j)].width = w
    an.column_dimensions["C"].width = 26

    # ---- 메모 시트
    memo = wb["메모"]
    memo.append([])
    memo.append(["주가수익률 (추가)"])
    memo.cell(memo.max_row, 1).font = BOLD
    memo.append(["가격 출처", "Yahoo Finance via yfinance 일별 종가"])
    memo.append(["기간", f"{s_lbl} 종가 → {e_lbl} 종가 (미국 장 마감 전이면 전 거래일 종가 사용)"])
    memo.append(["1년 주가수익률", "분할 조정 종가(Close) 기준, 배당 제외 = 종료 종가 / 시작 종가 - 1 (실적 시트 수식)"])
    memo.append(["1년 총수익률", "배당 재투자 조정 종가(Adj Close) 기준 = 종료 / 시작 - 1 (계산값)"])
    memo.append(["GICS 섹터", "Wikipedia List of S&P 500 companies"])
    if asof:
        memo.append(["1년 전 성장률·선정 기준", f"{os.path.basename(args.asof_file)} ({asof_lbl} 시점 SEC 제출 자료 기준)"])
    memo.append(["수익률 미산출 (시작일 이후 상장·분사)", ", ".join(no_start) or "없음"])

    wb.save(out)
    print(f"저장: {out}  | 기간 {s_lbl} → {e_lbl} | 수익률 미산출: {no_start}")


if __name__ == "__main__":
    sys.exit(main())
