#!/usr/bin/env python3
"""Russell 1000 종목 중 '최근 분기 매출이 회사 가이던스를 크게 상회한' 기업을 찾아 Excel로 정리한다.

단계:
    1) 종목·분기 실적: sp500_last5q.py 로 만든 Russell 1000 캐시(Yahoo 분기 매출·영업이익 5개 분기)
    2) 가이던스: 각 회사의 직전 분기 실적 발표문(SEC 8-K Item 2.02, Exhibit 99)에서
       최근 분기(Q0)에 대한 매출 가이던스를 추출 (guidance_sec.py)
    3) 비교: Q0 실제 매출 / 가이던스 중앙값 - 1, 가이던스 상단 초과 여부
    4) 상회 기업: PER·Forward PER·EPS 서프라이즈(Yahoo), 1년 주가수익률, 주요 사업 한 줄

사용법:
    python guidance_beats.py --fin-cache out/russell1000_cache_yahoo_YYYYMMDD.csv [--min-beat 0.05]

단계 2 결과는 out/r1000_guidance_YYYYMMDD.csv 에 캐시되어 중단 후 이어서 실행된다.
"""

import argparse
import csv
import os
import random
import sys
import time
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo

import numpy as np
import pandas as pd

import guidance_sec as gs
import sp500_last5q as base

ET = ZoneInfo("America/New_York")
KST = base.KST
GUIDE_FIELDS = ["ticker", "cik", "q0_end", "q0_release", "guide_filing", "guide_url", "low", "high", "mid", "kind",
                "flags", "snippet", "newer_release", "status"]


def load_fin(path):
    df = pd.read_csv(path, dtype={"ticker": str}).drop_duplicates("ticker", keep="last")
    return {r["ticker"]: r for r in df.to_dict("records")}


def find_guidance(ticker, cik, fin):
    """Q0 가이던스 추출. 반환 dict (GUIDE_FIELDS)."""
    out = {k: "" for k in GUIDE_FIELDS}
    out.update(ticker=ticker, cik=cik, status="no_guidance")
    q0, q1 = fin.get("date_0"), fin.get("date_1")
    rev0 = base._num(fin.get("rev_0"))
    if not isinstance(q0, str) or not isinstance(q1, str) or np.isnan(rev0):
        out["status"] = "no_financials"
        return out
    out["q0_end"] = q0
    ks = gs.earnings_8ks(cik)  # 최신순
    if not ks:
        out["status"] = "no_8k"
        return out
    after_q0 = [k for k in ks if k[0] > q0]
    if not after_q0:
        out["status"] = "q0_release_not_found"
        return out
    q0_rel = after_q0[-1][0]  # Q0 종료 후 첫 실적 발표
    out["q0_release"] = q0_rel
    out["newer_release"] = after_q0[0][0] if len(after_q0) > 1 and after_q0[0][0] > q0_rel else ""
    rev4 = base._num(fin.get("rev_4"))
    # Q-1 종료 후 ~ Q0 발표 전 실적 발표문 (최신 것부터 = 분기 중 가이던스 수정 반영)
    for date, acc, _ in [k for k in ks if q1 < k[0] < q0_rel]:
        texts = gs.exhibit_texts(cik, acc)
        res = gs.extract_revenue_guidance(texts, rev0, None if np.isnan(rev4) else rev4)
        if res:
            out.update(
                guide_filing=date,
                guide_url=gs.INDEX_HTM.format(cik=cik, acc=acc.replace("-", ""), acc_dash=acc),
                low=res["low"], high=res["high"], mid=res["mid"], kind=res["kind"], flags=res["flags"],
                snippet=res["snippet"][:500], status="ok",
            )
            return out
    return out


def collect_guidance(tickers, fin, cache_path):
    done = {}
    if os.path.exists(cache_path):
        done = {r["ticker"]: r for r in pd.read_csv(cache_path, dtype=str).fillna("").to_dict("records")}
    cmap = gs.cik_map()
    new = not os.path.exists(cache_path)
    with open(cache_path, "a", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=GUIDE_FIELDS)
        if new:
            w.writeheader()
        for i, t in enumerate(tickers, 1):
            if t in done:
                continue
            cik = cmap.get(t)
            if cik is None:
                rec = {k: "" for k in GUIDE_FIELDS}
                rec.update(ticker=t, status="no_cik")
            else:
                try:
                    rec = find_guidance(t, cik, fin.get(t, {}))
                except Exception as e:  # noqa: BLE001
                    rec = {k: "" for k in GUIDE_FIELDS}
                    rec.update(ticker=t, cik=cik, status=f"error: {type(e).__name__}")
            w.writerow(rec)
            f.flush()
            done[t] = rec
            if i % 25 == 0:
                ok = sum(1 for r in done.values() if r["status"] == "ok")
                print(f"  [{i}/{len(tickers)}] 가이던스 추출 {ok}개", flush=True)
    return done


def yahoo_extras(tickers):
    """PER, Forward PER, 섹터/업종, 최근 EPS 서프라이즈."""
    import yfinance as yf

    out = {}
    for t in tickers:
        d = {}
        for a in range(3):
            try:
                tk = yf.Ticker(t)
                info = tk.info
                d.update(pe=info.get("trailingPE"), fpe=info.get("forwardPE"), sector=info.get("sector"),
                         industry=info.get("industry"), name=info.get("shortName"))
                eh = tk.earnings_history
                if isinstance(eh, pd.DataFrame) and not eh.empty:
                    last = eh.sort_index().iloc[-1]
                    d.update(eps_act=last.get("epsActual"), eps_est=last.get("epsEstimate"),
                             eps_surp=last.get("surprisePercent"), eps_q=str(eh.sort_index().index[-1])[:10])
                break
            except Exception:  # noqa: BLE001
                time.sleep(2 * (a + 1))
        out[t] = d
        time.sleep(random.uniform(0.4, 0.8))
    return out


def price_returns(tickers, start):
    import yfinance as yf

    px = yf.download(tickers, start=(pd.Timestamp(start) - timedelta(days=7)).strftime("%Y-%m-%d"),
                     auto_adjust=False, group_by="column", progress=False, threads=True)["Close"]
    now = datetime.now(ET)
    if px.index.max().date() == now.date() and now.hour < 16:
        px = px.loc[px.index < px.index.max()]
    s = pd.Timestamp(start)
    end = px.index.max()
    ret = {}
    for t in tickers:
        if t in px and pd.notna(px.at[s, t]) and pd.notna(px.at[end, t]):
            ret[t] = (float(px.at[s, t]), float(px.at[end, t]), float(px.at[end, t] / px.at[s, t] - 1))
    return ret, s.strftime("%Y-%m-%d"), end.strftime("%Y-%m-%d")



def iwb_sectors():
    import requests
    from io import StringIO

    r = requests.get(base.IWB_HOLDINGS_CSV, headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"}, timeout=60)
    lines = r.text.splitlines()
    st = next(i for i, ln in enumerate(lines) if ln.startswith("Ticker,"))
    df = pd.read_csv(StringIO("\n".join(lines[st:])))
    df = df[df["Asset Class"] == "Equity"]
    return dict(zip(df["Ticker"].astype(str).str.strip().str.replace(r"[ .]", "-", regex=True), df["Sector"]))


def refresh_newer(guide, fin, gcache):
    """Q0 이후 새 실적 발표가 있었던 종목은 Yahoo 분기 실적을 다시 받아 가이던스 비교를 갱신."""
    cmap = None
    changed = 0
    for t, g in list(guide.items()):
        if not g.get("newer_release"):
            continue
        data, err = base.fetch_with_retry(base.fetch_yahoo, t)
        if not data or data["dates"][0] == fin.get(t, {}).get("date_0"):
            continue
        rec = {"ticker": t, "name": fin.get(t, {}).get("name", t), "status": "ok", "error": "", "note": data["note"]}
        for k in range(base.N_Q):
            rec[f"date_{k}"], rec[f"rev_{k}"], rec[f"oi_{k}"] = data["dates"][k], data["rev"][k], data["oi"][k]
        fin[t] = rec
        cmap = cmap or gs.cik_map()
        guide[t] = find_guidance(t, cmap[t], rec)
        changed += 1
    if changed:
        pd.DataFrame(list(guide.values()))[GUIDE_FIELDS].to_csv(gcache, index=False)
    print(f"새 분기 발표 반영: {changed}개")
    return changed


def build_excel(path, rows, allrows, sector_names, meta):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    HDR = PatternFill("solid", fgColor="DDEBF7")
    BOLD = Font(bold=True)
    M, P, X = "#,##0", "0.0%", "0.0"
    wb = Workbook()

    def header(ws, cols, row=1):
        for i, h in enumerate(cols, 1):
            c = ws.cell(row, i, h)
            c.font, c.fill = BOLD, HDR
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 45

    def ms(v):
        return None if v is None or (isinstance(v, float) and np.isnan(v)) else v / 1e6

    # ---------------- 1. 가이던스 상회 기업
    ws = wb.active
    ws.title = "가이던스 상회 기업"
    ql = base.Q_LABELS
    cols = (["기업명", "티커", "섹터(GICS)", "주요 사업", "최근 분기(Q0) 종료일", "Q0 실적 발표일",
             "매출 가이던스 하단 (백만$)", "매출 가이던스 상단 (백만$)", "가이던스 중앙값 (백만$)", "Q0 실제 매출 (백만$)",
             "가이던스 중앙값 대비", "가이던스 상단 대비", "가이던스 형태", "EPS 서프라이즈 (컨센서스 대비)"]
            + [f"매출 {q} (백만$)" for q in ql] + [f"영업이익 {q} (백만$)" for q in ql]
            + ["매출성장률(YoY)", "영업이익성장률(YoY)"] + [f"매출 QoQ({q})" for q in ql[:4]]
            + [f"주가 {meta['p_start']} ($)", f"주가 {meta['p_end']} ($)", "1년 주가수익률", "PER (trailing)", "Forward PER",
               "가이던스 근거 문장 (직전 분기 실적 발표문)", "SEC 공시"])
    header(ws, cols)
    C = {h: i + 1 for i, h in enumerate(cols)}
    L = {h: get_column_letter(i) for h, i in C.items()}
    for r, d in enumerate(rows, 2):
        v = {
            "기업명": d["name"], "티커": d["ticker"], "섹터(GICS)": d["sector"], "주요 사업": d["business"],
            "최근 분기(Q0) 종료일": d["q0_end"], "Q0 실적 발표일": d["q0_release"],
            "매출 가이던스 하단 (백만$)": ms(d["low"]), "매출 가이던스 상단 (백만$)": ms(d["high"]),
            "가이던스 중앙값 (백만$)": ms(d["mid"]), "Q0 실제 매출 (백만$)": ms(d["rev"][0]),
            "가이던스 중앙값 대비": f"={L['Q0 실제 매출 (백만$)']}{r}/{L['가이던스 중앙값 (백만$)']}{r}-1",
            "가이던스 상단 대비": f"={L['Q0 실제 매출 (백만$)']}{r}/{L['매출 가이던스 상단 (백만$)']}{r}-1",
            "가이던스 형태": d["kind_ko"], "EPS 서프라이즈 (컨센서스 대비)": d.get("eps_surp"),
            "영업이익성장률(YoY)": d["oi_yoy"],
            f"주가 {meta['p_start']} ($)": d.get("p0"), f"주가 {meta['p_end']} ($)": d.get("p1"),
            "PER (trailing)": d.get("pe"), "Forward PER": d.get("fpe"),
            "가이던스 근거 문장 (직전 분기 실적 발표문)": d["snippet"], "SEC 공시": d["guide_url"],
        }
        for k, q in enumerate(ql):
            v[f"매출 {q} (백만$)"] = ms(d["rev"][k])
            v[f"영업이익 {q} (백만$)"] = ms(d["oi"][k])
        r0, r4 = L[f"매출 {ql[0]} (백만$)"], L[f"매출 {ql[4]} (백만$)"]
        v["매출성장률(YoY)"] = f'=IF(AND(ISNUMBER({r0}{r}),ISNUMBER({r4}{r}),{r4}{r}<>0),{r0}{r}/{r4}{r}-1,"")'
        for k in range(4):
            a, b = L[f"매출 {ql[k]} (백만$)"], L[f"매출 {ql[k + 1]} (백만$)"]
            v[f"매출 QoQ({ql[k]})"] = f'=IF(AND(ISNUMBER({a}{r}),ISNUMBER({b}{r}),{b}{r}<>0),{a}{r}/{b}{r}-1,"")'
        a, b = L[f"주가 {meta['p_start']} ($)"], L[f"주가 {meta['p_end']} ($)"]
        v["1년 주가수익률"] = f'=IF(AND(ISNUMBER({a}{r}),ISNUMBER({b}{r})),{b}{r}/{a}{r}-1,"")'
        for h, val in v.items():
            if isinstance(val, float) and np.isnan(val):
                val = None
            c = ws.cell(r, C[h], val)
            if "(백만$)" in h:
                c.number_format = M
            elif "대비" in h or "성장률" in h or "QoQ" in h or "수익률" in h or "서프라이즈" in h:
                c.number_format = P
                if isinstance(val, str) and not val.startswith("="):
                    c.alignment = Alignment(horizontal="right")
            elif "PER" in h:
                c.number_format = X
            elif h.startswith("주가"):
                c.number_format = "#,##0.00"
            elif h == "SEC 공시" and val:
                c.hyperlink = val
                c.value = "8-K 보기"
                c.font = Font(color="0563C1", underline="single")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"
    widths = {"기업명": 24, "주요 사업": 44, "가이던스 근거 문장 (직전 분기 실적 발표문)": 70, "섹터(GICS)": 18}
    for h, i in C.items():
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 12)

    # ---------------- 2. 섹터 요약
    sv = wb.create_sheet("섹터 요약")
    header(sv, ["섹터(GICS)", "가이던스 비교 가능 기업 수", "크게 상회 기업 수", "크게 상회 비율",
                "가이던스 대비 중앙값 (섹터 전체)", "크게 상회 기업의 1년 주가수익률 중앙값", "크게 상회 기업"])
    full = "전체 가이던스 비교"
    n_all = len(allrows) + 1
    rng = lambda col: f"'{full}'!${col}$2:${col}${n_all}"  # noqa: E731
    beat_sheet_n = len(rows) + 1
    brng = lambda h: f"'가이던스 상회 기업'!${L[h]}$2:${L[h]}${beat_sheet_n}"  # noqa: E731
    from openpyxl.worksheet.formula import ArrayFormula

    for i, sname in enumerate(sector_names, 2):
        sv.cell(i, 1, sname)
        sv.cell(i, 2, f"=COUNTIF({rng('C')},A{i})")
        sv.cell(i, 3, f"=COUNTIF({brng('섹터(GICS)')},A{i})")
        sv.cell(i, 4, f'=IF(B{i}>0,C{i}/B{i},"")').number_format = P
        sv.cell(i, 5).value = ArrayFormula(f"E{i}", f'=IFERROR(MEDIAN(IF(({rng("C")}=A{i})*ISNUMBER({rng("J")}),{rng("J")})),"")')
        sv.cell(i, 5).number_format = P
        sv.cell(i, 6).value = ArrayFormula(
            f"F{i}", f'=IFERROR(MEDIAN(IF(({brng("섹터(GICS)")}=A{i})*ISNUMBER({brng("1년 주가수익률")}),{brng("1년 주가수익률")})),"")')
        sv.cell(i, 6).number_format = P
        sv.cell(i, 7, ", ".join(d["ticker"] for d in rows if d["sector"] == sname))
    t = len(sector_names) + 2
    sv.cell(t, 1, "합계").font = BOLD
    sv.cell(t, 2, f"=SUM(B2:B{t - 1})")
    sv.cell(t, 3, f"=SUM(C2:C{t - 1})")
    sv.cell(t, 4, f'=IF(B{t}>0,C{t}/B{t},"")').number_format = P
    for col, w in zip("ABCDEFG", (24, 14, 12, 12, 16, 18, 90)):
        sv.column_dimensions[col].width = w

    # ---------------- 3. 전체 가이던스 비교
    fa = wb.create_sheet(full)
    fcols = ["기업명", "티커", "섹터(GICS)", "Q0 종료일", "가이던스 하단 (백만$)", "가이던스 상단 (백만$)",
             "가이던스 중앙값 (백만$)", "Q0 실제 매출 (백만$)", "가이던스 형태", "가이던스 중앙값 대비", "가이던스 상단 대비",
             "크게 상회", "매출성장률(YoY)", "가이던스 근거 문장", "SEC 공시"]
    header(fa, fcols)
    for r, d in enumerate(allrows, 2):
        vals = [d["name"], d["ticker"], d["sector"], d["q0_end"], ms(d["low"]), ms(d["high"]), ms(d["mid"]),
                ms(d["rev"][0]), d["kind_ko"], f"=H{r}/G{r}-1", f"=H{r}/F{r}-1",
                f'=IF(AND(J{r}>={meta["min_beat"]},K{r}>0),"예","")',
                (d["rev"][0] / d["rev"][4] - 1) if d["rev"][4] and not np.isnan(d["rev"][4]) and d["rev"][4] != 0 else None,
                d["snippet"], d["guide_url"]]
        for i, val in enumerate(vals, 1):
            if isinstance(val, float) and np.isnan(val):
                val = None
            c = fa.cell(r, i, val)
            if 5 <= i <= 8:
                c.number_format = M
            elif i in (10, 11, 13):
                c.number_format = P
            elif i == 15 and val:
                c.hyperlink, c.value, c.font = val, "8-K 보기", Font(color="0563C1", underline="single")
    fa.freeze_panes = "C2"
    fa.auto_filter.ref = f"A1:O{len(allrows) + 1}"
    for i, w in enumerate([24, 8, 18, 11, 12, 12, 12, 12, 12, 11, 11, 8, 11, 80, 10], 1):
        fa.column_dimensions[get_column_letter(i)].width = w

    # ---------------- 4. 메모
    mm = wb.create_sheet("메모")
    notes = [
        ("작성 기준일", meta["asof"]),
        ("대상", f"Russell 1000 (iShares IWB 보유종목, {meta['n_universe']}개 기업, 동일 기업 복수 주식 클래스 제외)"),
        ("최근 분기(Q0)", "각 기업이 가장 최근에 실적을 발표한 분기 (Yahoo Finance 분기 손익계산서 기준, 값이 빈 최신 분기는 제외)"),
        ("회사 가이던스", "Q0 직전 분기 실적 발표문(SEC 8-K Item 2.02, Exhibit 99)에 회사가 제시한 'Q0 매출 전망'. "
                       "분기 중 수정 발표가 있으면 가장 최근 것. 규칙 기반 문장 추출 → 근거 문장 열에 원문 표기"),
        ("가이던스 비교 가능", f"{meta['n_guided']}개 기업 (분기 매출 가이던스를 숫자로 제시한 기업). "
                         "나머지는 가이던스 미제공·연간 가이던스만 제공·EPS만 제공·보도자료 외(컨퍼런스콜)에서만 제공"),
        ("'크게 상회' 기준", f"Q0 실제 매출이 가이던스 중앙값보다 {meta['min_beat']:.0%} 이상 높고, 가이던스 상단도 넘은 경우"),
        ("가이던스 형태", "범위(하단~상단), ±금액, ±%, 단일값(approximately), 성장률(전년 동기 매출 × (1+가이던스 성장률)로 환산)"),
        ("EPS 서프라이즈", "Yahoo Finance: 최근 분기 실제 EPS vs 애널리스트 컨센서스 (회사 가이던스가 아닌 참고 지표)"),
        ("주가수익률", f"{meta['p_start']} 종가 → {meta['p_end']} 종가 (Yahoo, 분할 조정·배당 제외)"),
        ("PER / Forward PER", "Yahoo Finance trailingPE / forwardPE (수집 시점 주가 기준, 적자 기업은 빈칸)"),
        ("매출·영업이익", "Yahoo Finance quarterly_income_stmt, 백만 달러. 영업이익 YoY 규칙: Q-4 ≤ 0 → 흑자전환/N/M, Q-4>0 & Q0<0 → 적자전환"),
        ("주의", "가이던스와 실제 매출의 기준(GAAP/비GAAP, 환율·인수 효과 등)이 다를 수 있음. 인수합병으로 매출이 늘어난 경우 "
               "가이던스에 반영되지 않았을 수 있으니 근거 문장을 확인할 것. 투자 판단의 근거가 아닌 스크리닝 결과임"),
    ]
    for i, (a, b) in enumerate(notes, 1):
        mm.cell(i, 1, a).font = BOLD
        mm.cell(i, 2, b)
    mm.column_dimensions["A"].width = 20
    mm.column_dimensions["B"].width = 140
    wb.save(path)


KIND_KO = {"range": "범위", "±$": "±금액", "±%": "±%", "point": "단일값", "growth%": "성장률(환산)"}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--fin-cache", required=True, help="Russell 1000 분기 실적 캐시 CSV (sp500_last5q.py 결과)")
    ap.add_argument("--min-beat", type=float, default=0.05, help="'크게 상회' 기준: 가이던스 중앙값 대비 (기본 5%%)")
    ap.add_argument("--outdir", default="out")
    ap.add_argument("--stage", choices=["guidance", "all"], default="all")
    args = ap.parse_args()

    today = datetime.now(KST).strftime("%Y%m%d")
    universe, uni_src = base.get_russell1000_list()
    fin = load_fin(args.fin_cache)
    universe, dups = base.dedupe_share_classes(universe, fin)
    tickers = list(universe["ticker"])
    print(f"대상 {len(tickers)}개 (중복 클래스 {len(dups)}개 제외)")

    gs.TEXT_CACHE_DIR = os.path.join(args.outdir, "sec_text_cache")
    gcache = os.path.join(args.outdir, f"r1000_guidance_{today}.csv")
    guide = collect_guidance(tickers, fin, gcache)
    print("가이던스 단계 완료:", pd.Series([r["status"] for r in guide.values()]).value_counts().to_dict())
    if args.stage == "guidance":
        return 0

    refresh_newer(guide, fin, gcache)

    # 수동 검토 결과: guidance_overrides.csv (ticker, action[exclude|set], low, high, reason)
    ov_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "guidance_overrides.csv")
    overrides = {}
    if os.path.exists(ov_path):
        overrides = {r["ticker"]: r for r in pd.read_csv(ov_path, dtype=str).fillna("").to_dict("records")}

    sectors = iwb_sectors()
    names = dict(zip(universe["ticker"], universe["name"]))
    allrows = []
    for t in tickers:
        g = guide.get(t, {})
        ov = overrides.get(t)
        if ov and ov["action"] == "exclude":
            continue
        if g.get("status") != "ok" and not (ov and ov["action"] == "set"):
            continue
        f = fin.get(t, {})
        d = {
            "ticker": t, "name": names.get(t, t), "sector": sectors.get(t, ""),
            "q0_end": g.get("q0_end", ""), "q0_release": g.get("q0_release", ""),
            "low": float(g["low"]) if g.get("low") else np.nan, "high": float(g["high"]) if g.get("high") else np.nan,
            "kind": g.get("kind", ""), "snippet": g.get("snippet", ""), "guide_url": g.get("guide_url", ""),
            "rev": [base._num(f.get(f"rev_{k}")) for k in range(base.N_Q)],
            "oi": [base._num(f.get(f"oi_{k}")) for k in range(base.N_Q)],
        }
        if ov and ov["action"] == "set":
            d["low"], d["high"] = float(ov["low"]), float(ov["high"])
            d["snippet"] = f"[수동 확인] {ov['reason']} | " + d["snippet"]
        d["mid"] = (d["low"] + d["high"]) / 2
        d["kind_ko"] = KIND_KO.get(d["kind"], d["kind"])
        d["beat"] = d["rev"][0] / d["mid"] - 1
        d["beat_hi"] = d["rev"][0] / d["high"] - 1
        d["oi_yoy"] = base.op_income_yoy(d["oi"][0], d["oi"][4])
        allrows.append(d)
    allrows.sort(key=lambda d: -d["beat"])
    rows = [d for d in allrows if d["beat"] >= args.min_beat and d["beat_hi"] > 0]
    print(f"가이던스 비교 {len(allrows)}개, 크게 상회 {len(rows)}개")

    beat_t = [d["ticker"] for d in rows]
    start = (datetime.now(ET).date() - timedelta(days=365))
    while start.weekday() >= 5:
        start -= timedelta(days=1)
    rets, p_start, p_end = price_returns(beat_t, start.strftime("%Y-%m-%d"))
    extras = yahoo_extras(beat_t)
    ko = base.load_business_ko()
    for d in rows:
        e = extras.get(d["ticker"], {})
        d.update(pe=e.get("pe"), fpe=e.get("fpe"), eps_surp=e.get("eps_surp"))
        if d["ticker"] in rets:
            d["p0"], d["p1"], _ = rets[d["ticker"]]
        ind = " / ".join(x for x in (e.get("sector"), e.get("industry")) if x)
        d["business"] = ko.get(d["ticker"]) or (f"(Yahoo 업종) {ind}" if ind else "")
        if e.get("name") and d["name"].isupper():
            d["name"] = e["name"]
    meta = {"asof": datetime.now(KST).strftime("%Y-%m-%d %H:%M KST"), "n_universe": len(tickers), "n_guided": len(allrows),
            "min_beat": args.min_beat, "p_start": p_start, "p_end": p_end}
    sector_names = sorted({d["sector"] for d in allrows if d["sector"]})
    out = os.path.join(args.outdir, f"russell1000_guidance_beats_{today}.xlsx")
    build_excel(out, rows, allrows, sector_names, meta)
    print("저장:", out)
    print("설명 없는 종목:", [d["ticker"] for d in rows if not d["business"] or d["business"].startswith("(Yahoo")])
    return 0


if __name__ == "__main__":
    sys.exit(main())
