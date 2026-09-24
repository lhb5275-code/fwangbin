#!/usr/bin/env python3
"""S&P 500 구성 종목의 최근 5개 분기 매출/영업이익을 Yahoo Finance(yfinance)에서 수집해 Excel로 정리한다.

사용법:
    python sp500_last5q.py                 # 전체 S&P 500 수집 (캐시가 있으면 이어서)
    python sp500_last5q.py --tickers AAPL MSFT NVDA JPM   # 일부 종목만 (테스트용)
    python sp500_last5q.py --fresh         # 캐시 무시하고 처음부터

중간 결과는 sp500_cache_YYYYMMDD.csv 에 종목별로 한 줄씩 추가된다.
중단 후 다시 실행하면 이미 성공(ok/partial)한 종목은 건너뛰고, 실패한 종목만 다시 시도한다.
"""

import argparse
import csv
import os
import random
import sys
import time
import unicodedata
from datetime import datetime, timezone, timedelta
from io import StringIO

import numpy as np
import pandas as pd
import requests
import yfinance as yf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

WIKI_URL = "https://en.wikipedia.org/wiki/List_of_S%26P_500_companies"
N_Q = 5
MAX_RETRIES = 3
Q_LABELS = ["Q0", "Q-1", "Q-2", "Q-3", "Q-4"]
VERIFY_TICKERS = ["AAPL", "MSFT", "NVDA", "JPM"]
KST = timezone(timedelta(hours=9))

CACHE_FIELDS = (
    ["ticker", "name", "status", "error"]
    + [f"date_{i}" for i in range(N_Q)]
    + [f"rev_{i}" for i in range(N_Q)]
    + [f"oi_{i}" for i in range(N_Q)]
)


# --------------------------------------------------------------------------- 1. 종목 리스트
def get_sp500_list():
    """Wikipedia 표에서 (기업명, Yahoo 티커) 리스트를 가져온다."""
    headers = {"User-Agent": "Mozilla/5.0 (sp500-last5q script)"}
    resp = requests.get(WIKI_URL, headers=headers, timeout=30)
    resp.raise_for_status()
    tables = pd.read_html(StringIO(resp.text), attrs={"id": "constituents"})
    df = tables[0]
    out = pd.DataFrame(
        {
            "name": df["Security"].astype(str).str.strip(),
            "ticker": df["Symbol"].astype(str).str.strip().str.replace(".", "-", regex=False),
        }
    )
    return out.drop_duplicates("ticker").reset_index(drop=True)


# --------------------------------------------------------------------------- 2. 데이터 수집
def _row(stmt, label):
    return stmt.loc[label] if label in stmt.index else None


def fetch_quarters(ticker):
    """yfinance에서 최근 5개 분기 매출/영업이익을 가져온다.

    반환: dict(dates, rev, oi) - 각 길이 N_Q, 없는 값은 NaN/None.
    데이터가 비어 있으면 예외를 던져 재시도 대상으로 만든다.
    """
    stmt = yf.Ticker(ticker).quarterly_income_stmt
    if stmt is None or stmt.empty:
        raise RuntimeError("quarterly_income_stmt 가 비어 있음")

    cols = sorted(stmt.columns, key=pd.Timestamp, reverse=True)[:N_Q]
    rev_row = _row(stmt, "Total Revenue")
    oi_row = _row(stmt, "Operating Income")

    dates, rev, oi = [], [], []
    for i in range(N_Q):
        if i < len(cols):
            c = cols[i]
            dates.append(pd.Timestamp(c).strftime("%Y-%m-%d"))
            rev.append(float(rev_row[c]) if rev_row is not None and pd.notna(rev_row[c]) else np.nan)
            oi.append(float(oi_row[c]) if oi_row is not None and pd.notna(oi_row[c]) else np.nan)
        else:
            dates.append(None)
            rev.append(np.nan)
            oi.append(np.nan)
    return {"dates": dates, "rev": rev, "oi": oi, "has_oi_row": oi_row is not None}


def polite_sleep():
    time.sleep(random.uniform(0.5, 1.0))


def fetch_with_retry(ticker):
    last_err = None
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            return fetch_quarters(ticker), None
        except Exception as e:  # noqa: BLE001 - 네트워크/파싱 오류 모두 재시도
            last_err = f"{type(e).__name__}: {e}"
            print(f"    [{ticker}] 시도 {attempt}/{MAX_RETRIES} 실패 - {last_err}")
            if attempt < MAX_RETRIES:
                time.sleep(attempt * random.uniform(1.0, 2.0))
        finally:
            polite_sleep()
    return None, last_err


def load_cache(path):
    if not os.path.exists(path):
        return {}
    df = pd.read_csv(path, dtype={"ticker": str})
    # 같은 티커가 여러 번 기록되었으면(재시도) 마지막 기록을 사용
    df = df.drop_duplicates("ticker", keep="last")
    return {r["ticker"]: r for r in df.to_dict("records")}


def append_cache(path, rec):
    new = not os.path.exists(path)
    with open(path, "a", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=CACHE_FIELDS)
        if new:
            w.writeheader()
        w.writerow(rec)


def collect(universe, cache_path):
    cache = load_cache(cache_path)
    total = len(universe)
    for i, (name, ticker) in enumerate(zip(universe["name"], universe["ticker"]), 1):
        cached = cache.get(ticker)
        if cached is not None and cached["status"] in ("ok", "partial"):
            continue
        print(f"[{i}/{total}] {ticker} ({name})")
        data, err = fetch_with_retry(ticker)
        rec = {"ticker": ticker, "name": name}
        if data is None:
            rec.update(status="failed", error=err)
        else:
            complete = all(d is not None for d in data["dates"]) and not any(
                np.isnan(v) for v in data["rev"] + data["oi"]
            )
            rec.update(status="ok" if complete else "partial", error="")
            for k in range(N_Q):
                rec[f"date_{k}"] = data["dates"][k]
                rec[f"rev_{k}"] = data["rev"][k]
                rec[f"oi_{k}"] = data["oi"][k]
        append_cache(cache_path, rec)
        cache[ticker] = rec
    return cache


# --------------------------------------------------------------------------- 3. 계산
def _num(v):
    try:
        v = float(v)
    except (TypeError, ValueError):
        return np.nan
    return v


def ratio_growth(cur, base):
    cur, base = _num(cur), _num(base)
    if np.isnan(cur) or np.isnan(base) or base == 0:
        return np.nan
    return cur / base - 1


def op_income_yoy(q0, q4):
    """영업이익 YoY. 기준(Q-4)이 0 이하이면 N/M, 부호 전환은 텍스트로 표시."""
    q0, q4 = _num(q0), _num(q4)
    if np.isnan(q0) or np.isnan(q4):
        return np.nan
    if q4 <= 0:
        return "흑자전환" if q0 > 0 else "N/M"
    if q0 < 0:
        return "적자전환"
    return q0 / q4 - 1


def build_table(universe, cache):
    rows = []
    for name, ticker in zip(universe["name"], universe["ticker"]):
        r = cache.get(ticker) or {}
        rev = [_num(r.get(f"rev_{k}")) for k in range(N_Q)]
        oi = [_num(r.get(f"oi_{k}")) for k in range(N_Q)]
        dates = [r.get(f"date_{k}") if isinstance(r.get(f"date_{k}"), str) else None for k in range(N_Q)]
        row = {
            "name": name,
            "ticker": ticker,
            "status": r.get("status", "failed"),
            "error": r.get("error") if isinstance(r.get("error"), str) else "",
            "dates": dates,
            "rev": rev,
            "oi": oi,
            "rev_yoy": ratio_growth(rev[0], rev[4]),
            "oi_yoy": op_income_yoy(oi[0], oi[4]),
            "qoq": [ratio_growth(rev[k], rev[k + 1]) for k in range(4)],
        }
        rows.append(row)
    # 매출 YoY 내림차순, NaN은 맨 아래
    rows.sort(key=lambda x: (np.isnan(x["rev_yoy"]), -x["rev_yoy"] if not np.isnan(x["rev_yoy"]) else 0))
    return rows


# --------------------------------------------------------------------------- 4. Excel
HEADERS = (
    ["기업명", "티커", "최근 분기 종료일"]
    + [f"매출 {q} (백만$)" for q in Q_LABELS]
    + [f"영업이익 {q} (백만$)" for q in Q_LABELS]
    + ["매출성장률(YoY)", "영업이익성장률(YoY)"]
    + [f"매출 QoQ({q})" for q in Q_LABELS[:4]]
)
FMT_MUSD = "#,##0"
FMT_PCT = "0.0%"


def _display_width(s):
    return sum(2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1 for ch in str(s))


def _cell_text(v, fmt):
    if v is None:
        return ""
    if isinstance(v, float):
        if fmt == FMT_PCT:
            return f"{v * 100:.1f}%"
        if fmt == FMT_MUSD:
            return f"{v:,.0f}"
    return str(v)


def _blank(v):
    return None if isinstance(v, float) and np.isnan(v) else v


def _musd(v):
    return None if np.isnan(v) else v / 1e6


def autofit(ws, formats, min_w=6, max_w=60):
    widths = {}
    for row in ws.iter_rows():
        for c in row:
            w = _display_width(_cell_text(c.value, formats.get(c.column)))
            widths[c.column] = max(widths.get(c.column, 0), w)
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max(w + 2, min_w), max_w)


def write_excel(rows, universe_size, failed, missing, collected_at, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "실적"
    ws.append(HEADERS)
    hdr_fill = PatternFill("solid", fgColor="DDEBF7")
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)

    for r in rows:
        ws.append(
            [r["name"], r["ticker"], r["dates"][0]]
            + [_musd(v) for v in r["rev"]]
            + [_musd(v) for v in r["oi"]]
            + [_blank(r["rev_yoy"]), _blank(r["oi_yoy"])]
            + [_blank(v) for v in r["qoq"]]
        )

    formats = {}
    for col in range(4, 14):
        formats[col] = FMT_MUSD
    for col in range(14, 20):
        formats[col] = FMT_PCT
    for row in ws.iter_rows(min_row=2):
        for c in row:
            fmt = formats.get(c.column)
            if fmt and isinstance(c.value, (int, float)):
                c.number_format = fmt
            elif c.column == 15 and isinstance(c.value, str):
                c.alignment = Alignment(horizontal="right")

    ws.freeze_panes = "C2"  # 헤더 행 + 기업명/티커 열 고정
    ws.auto_filter.ref = ws.dimensions
    autofit(ws, formats)

    # ---- 메모 시트
    memo = wb.create_sheet("메모")
    bold = Font(bold=True)
    ok_count = sum(1 for r in rows if r["status"] in ("ok", "partial"))
    memo.append(["데이터 출처", "Yahoo Finance via yfinance (yf.Ticker(t).quarterly_income_stmt)"])
    memo.append(["종목 리스트 출처", WIKI_URL])
    memo.append(["수집 일시", collected_at])
    memo.append(["수집 성공 / 전체", f"{ok_count} / {universe_size}"])
    memo.append(["단위", "매출·영업이익: 백만 달러(USD mn), 성장률: %"])
    memo.append(
        [
            "영업이익 YoY 규칙",
            "Q-4 영업이익 ≤ 0 → Q0 > 0 이면 '흑자전환', 아니면 'N/M' / Q-4 > 0 이고 Q0 < 0 → '적자전환'",
        ]
    )
    for r in range(1, memo.max_row + 1):
        memo.cell(r, 1).font = bold

    memo.append([])
    memo.append([f"수집 실패 티커 ({len(failed)}개)"])
    memo.cell(memo.max_row, 1).font = bold
    memo.append(["티커", "기업명", "오류"])
    for c in memo[memo.max_row]:
        c.font = bold
    for t, n, e in failed:
        memo.append([t, n, e])
    if not failed:
        memo.append(["(없음)"])

    memo.append([])
    memo.append([f"데이터 누락 티커 ({len(missing)}개)"])
    memo.cell(memo.max_row, 1).font = bold
    memo.append(["티커", "기업명", "누락 내용"])
    for c in memo[memo.max_row]:
        c.font = bold
    for t, n, e in missing:
        memo.append([t, n, e])
    if not missing:
        memo.append(["(없음)"])

    memo.append([])
    memo.append(["기업별 실제 분기 종료일"])
    memo.cell(memo.max_row, 1).font = bold
    memo.append(["티커", "기업명"] + [f"{q} 종료일" for q in Q_LABELS])
    for c in memo[memo.max_row]:
        c.font = bold
    for r in sorted(rows, key=lambda x: x["ticker"]):
        memo.append([r["ticker"], r["name"]] + [d or "" for d in r["dates"]])
    autofit(memo, {}, max_w=80)

    wb.save(path)


def describe_missing(r):
    if r["status"] == "failed":
        return None
    issues = []
    n_dates = sum(1 for d in r["dates"] if d)
    if n_dates < N_Q:
        issues.append(f"분기 {n_dates}개만 제공")
    miss_rev = [Q_LABELS[k] for k in range(n_dates) if np.isnan(r["rev"][k])]
    miss_oi = [Q_LABELS[k] for k in range(n_dates) if np.isnan(r["oi"][k])]
    if miss_rev:
        issues.append("매출 없음: " + ",".join(miss_rev))
    if miss_oi:
        issues.append("영업이익 없음: " + ",".join(miss_oi))
    if r["dates"][0] and r["dates"][4]:
        gap = (pd.Timestamp(r["dates"][0]) - pd.Timestamp(r["dates"][4])).days
        if not 330 <= gap <= 400:
            issues.append(f"Q0~Q-4 간격 {gap}일(1년 아님, YoY 해석 주의)")
    return "; ".join(issues) or None


# --------------------------------------------------------------------------- 5. 검증 출력
def _fmt_m(v):
    return "" if np.isnan(v) else f"{v / 1e6:,.0f}"


def _fmt_p(v):
    if isinstance(v, str):
        return v
    return "" if np.isnan(v) else f"{v * 100:.1f}%"


def print_verification(rows):
    by_t = {r["ticker"]: r for r in rows}
    print("\n" + "=" * 78)
    print("검증용 값 (단위: 백만$) - Yahoo Finance > Financials > Income Statement > Quarterly 와 대조")
    print("=" * 78)
    for t in VERIFY_TICKERS:
        r = by_t.get(t)
        if r is None:
            print(f"\n{t}: 대상 목록에 없음")
            continue
        print(f"\n{t} - {r['name']}  (status: {r['status']})")
        print(f"  {'분기':<5}{'종료일':<12}{'매출':>14}{'영업이익':>14}")
        for k in range(N_Q):
            print(f"  {Q_LABELS[k]:<6}{r['dates'][k] or '-':<14}{_fmt_m(r['rev'][k]):>14}{_fmt_m(r['oi'][k]):>16}")
        print(f"  매출 YoY: {_fmt_p(r['rev_yoy'])}   영업이익 YoY: {_fmt_p(r['oi_yoy'])}")
        print("  매출 QoQ: " + ", ".join(f"{Q_LABELS[k]} {_fmt_p(r['qoq'][k])}" for k in range(4)))


# --------------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--tickers", nargs="+", help="전체 대신 이 티커들만 수집 (Yahoo 형식, 예: BRK-B)")
    ap.add_argument("--fresh", action="store_true", help="기존 캐시 CSV를 지우고 처음부터 수집")
    ap.add_argument("--outdir", default=".", help="출력/캐시 디렉터리 (기본: 현재 디렉터리)")
    args = ap.parse_args()

    today = datetime.now(KST).strftime("%Y%m%d")
    os.makedirs(args.outdir, exist_ok=True)
    cache_path = os.path.join(args.outdir, f"sp500_cache_{today}.csv")
    out_path = os.path.join(args.outdir, f"sp500_last5q_{today}.xlsx")
    if args.fresh and os.path.exists(cache_path):
        os.remove(cache_path)

    print("S&P 500 종목 리스트 가져오는 중 (Wikipedia)...")
    universe = get_sp500_list()
    if args.tickers:
        wanted = [t.upper().replace(".", "-") for t in args.tickers]
        sub = universe[universe["ticker"].isin(wanted)]
        extra = [t for t in wanted if t not in set(sub["ticker"])]
        universe = pd.concat([sub, pd.DataFrame({"name": extra, "ticker": extra})], ignore_index=True)
    print(f"대상 종목: {len(universe)}개  |  캐시: {cache_path}")

    cache = collect(universe, cache_path)
    collected_at = datetime.now(KST).strftime("%Y-%m-%d %H:%M:%S KST")

    rows = build_table(universe, cache)
    failed = [(r["ticker"], r["name"], r["error"]) for r in rows if r["status"] == "failed"]
    missing = [(r["ticker"], r["name"], m) for r in rows if (m := describe_missing(r))]
    write_excel(rows, len(universe), failed, missing, collected_at, out_path)

    print_verification(rows)

    ok = len(universe) - len(failed)
    print("\n" + "=" * 78)
    print(f"수집 성공: {ok} / {len(universe)} 종목  (데이터 일부 누락: {len(missing)}개)")
    print(f"최종 실패 티커 ({len(failed)}개): {', '.join(t for t, _, _ in failed) or '없음'}")
    print(f"Excel 저장: {out_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
