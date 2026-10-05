#!/usr/bin/env python3
"""기업분석 PPT 생성기.

companies/<TICKER>/deck.yaml 원고를 읽어 '구글 기업분석' 예시와 같은 양식의 .pptx를 만든다.

    python build_deck.py companies/GOOGL_sample/deck.yaml
    python build_deck.py companies/NVDA/deck.yaml -o output/NVDA_기업분석.pptx

양식 요약 (STYLE.md 참고)
- 16:9 (13.333" x 7.5"), 검정 배경
- 왼쪽 위 흰색 제목 탭: 회사 로고 + 제목 (에스코어 드림 7 ExtraBold 24pt, 검정)
- 본문: 에스코어 드림 6 Bold 20pt, 흰색. [[노랑 강조]] {{연두 강조}} ((살구 강조))
- 표지/구역/감사 슬라이드: 여기어때 잘난체 고딕
- 슬라이드마다 발표 대본(노트), 화면 전환 효과 없음
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

import yaml
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# ---------------------------------------------------------------- 스타일 상수
SLIDE_W, SLIDE_H = 13.333, 7.5

FONT_BODY = "에스코어 드림 6 Bold"
FONT_TITLE = "에스코어 드림 7 ExtraBold"
FONT_DISPLAY = "여기어때 잘난체 고딕"
FONT_LIGHT = "에스코어 드림 4 Regular"   # 기준일 표기 등 보조 글자

BLACK = "000000"
WHITE = "FFFFFF"
YELLOW = "FFFF00"       # 핵심 수치·키워드
GREEN = "B4E5A2"        # 증감률·보조 강조 (Office accent6, 밝기 +60%)
PEACH = "F6C6AD"        # 전망치·미래 구간 (Office accent2, 밝기 +60%)
SKY = "CAEEFB"          # 막대그래프
LIME = "D9E550"         # 사업 구조도 상자
TABLE_HEAD = "DAF2D0"   # 표 머리글 (Office accent6, 밝기 +80%)
GRAY = "9AA0A6"
DARK = "202124"
UP_RED = "D93025"       # 주가 상승선 (국내 관례: 상승=빨강)
METRIC_RED = "FF0000"   # 가이던스 차트: 영업이익률
METRIC_GREEN = "00B050" # 가이던스 차트: EPS
DOWN_BLUE = "1A73E8"

TAB_Y, TAB_H = 0.30, 0.47
LOGO_X, LOGO_Y, LOGO_H, LOGO_MAX_W = 0.13, 0.36, 0.34, 1.6
BODY_X, BODY_Y, BODY_W = 0.92, 1.02, 11.5
BODY_PT = 20
LINE_IN = BODY_PT * 1.2 / 72 + 4 / 72      # 한 줄 높이 추정 (줄간격 + 단락 앞 간격)
CONTENT_BOTTOM = 7.05

MARK_RE = re.compile(r"(\[\[.*?\]\]|\{\{.*?\}\}|\(\(.*?\)\))")


# ---------------------------------------------------------------- 유틸
def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_.upper())


def char_em(ch: str) -> float:
    o = ord(ch)
    if 0xAC00 <= o <= 0xD7A3 or 0x3130 <= o <= 0x318F or 0x4E00 <= o <= 0x9FFF:
        return 1.0
    if ch == " ":
        return 0.3
    if ch.isupper() or ch.isdigit() or ch in "$%&@#":
        return 0.66
    if ch.isalpha():
        return 0.56
    return 0.38


def text_width(text: str, pt: float) -> float:
    """대략적인 글자 폭(인치). 에스코어 드림 기준으로 약간 넉넉하게 잡는다."""
    plain = MARK_RE.sub(lambda m: m.group(0)[2:-2], text)
    return sum(char_em(c) for c in plain) * pt / 72 * 1.02


def wrapped_lines(text: str, pt: float, width: float) -> int:
    if not text.strip():
        return 1
    return max(1, int(text_width(text, pt) / max(width, 0.5)) + 1)


def set_run_font(run, face: str, size: float | None = None, color: str | None = None,
                 bold: bool | None = None, highlight: str | None = None):
    f = run.font
    if size:
        f.size = Pt(size)
    if color:
        f.color.rgb = rgb(color)
    if bold is not None:
        f.bold = bold
    f.name = face
    rPr = run._r.get_or_add_rPr()
    latin = rPr.find(qn("a:latin"))
    for tag in ("a:ea", "a:cs"):
        old = rPr.find(qn(tag))
        if old is not None:
            rPr.remove(old)
    ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", face)
    latin.addnext(ea)
    if highlight:
        hl = etree.Element(qn("a:highlight"))
        clr = etree.SubElement(hl, qn("a:srgbClr"))
        clr.set("val", highlight)
        latin.addprevious(hl)
    rPr.set("lang", "ko-KR")
    rPr.set("altLang", "en-US")


def add_rich_text(paragraph, text: str, size: float, base_color: str = WHITE,
                  face: str = FONT_BODY, on_light: bool = False):
    """[[노랑]] {{연두}} ((살구)) 표기를 색 있는 런으로 바꿔 넣는다.

    밝은 배경(on_light) 위에서는 노랑을 형광펜 배경으로, 연두·살구를 진한 색 글자로 쓴다.
    """
    for tok in MARK_RE.split(text):
        if not tok:
            continue
        color, hl = base_color, None
        if tok.startswith("[["):
            tok = tok[2:-2]
            color, hl = (base_color, YELLOW) if on_light else (YELLOW, None)
        elif tok.startswith("{{"):
            tok = tok[2:-2]
            color = "188038" if on_light else GREEN
        elif tok.startswith("(("):
            tok = tok[2:-2]
            color = "C5221F" if on_light else PEACH
        r = paragraph.add_run()
        r.text = tok
        set_run_font(r, face, size, color, highlight=hl)


def set_bullet(paragraph, char: str | None, mar_l: float, indent: float):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(mar_l))))
    pPr.set("indent", str(int(Inches(indent))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buFont"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if char:
        bf = etree.SubElement(pPr, qn("a:buFont"))
        bf.set("typeface", "Arial")
        bc = etree.SubElement(pPr, qn("a:buChar"))
        bc.set("char", char)
    else:
        etree.SubElement(pPr, qn("a:buNone"))


def para_spacing(paragraph, before_pt: float = 4, line: float = 1.0):
    paragraph.space_before = Pt(before_pt)
    paragraph.line_spacing = line


def text_box(slide, x, y, w, h, *, name=None, anchor=MSO_ANCHOR.TOP, wrap=True, margin=0.05):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    return tb, tf


def no_line(shape):
    shape.line.fill.background()


def solid(shape, hex_):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(hex_)


def set_background(slide, hex_):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb(hex_)


def add_glow(run, color: str = BLACK, rad_pt: float = 10):
    """사진 위 글자가 잘 보이도록 글자 주위에 번지는 그림자(광선) 효과를 준다."""
    rPr = run._r.get_or_add_rPr()
    eff = etree.Element(qn("a:effectLst"))
    glow = etree.SubElement(eff, qn("a:glow"))
    glow.set("rad", str(int(rad_pt * 12700)))
    c = etree.SubElement(glow, qn("a:srgbClr"))
    c.set("val", color)
    anchor = rPr.find(qn("a:highlight"))
    if anchor is None:
        anchor = rPr.find(qn("a:latin"))
    anchor.addprevious(eff)


def add_fill_picture(slide, path: Path, x: float, y: float, w: float, h: float, name: str | None = None):
    """상자를 꽉 채우도록 사진을 넣고, 비율이 안 맞는 부분은 잘라낸다 (늘이지 않음)."""
    pw, ph = image_size(path)
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    box_ar, img_ar = w / h, pw / ph
    if img_ar > box_ar:          # 사진이 더 넓다 → 좌우를 자름
        cut = (1 - box_ar / img_ar) / 2
        pic.crop_left = pic.crop_right = cut
    elif img_ar < box_ar:        # 사진이 더 길다 → 위아래를 자름 (위쪽은 덜 잘라 얼굴·간판 보존)
        cut = 1 - img_ar / box_ar
        pic.crop_top, pic.crop_bottom = cut * 0.3, cut * 0.7
    if name:
        pic.name = name
    return pic


def image_size(path: Path) -> tuple[int, int]:
    with Image.open(path) as im:
        return im.size


def fit(path: Path, max_w: float, max_h: float) -> tuple[float, float]:
    pw, ph = image_size(path)
    s = min(max_w / pw, max_h / ph)
    return pw * s, ph * s


# ---------------------------------------------------------------- 빌더
class DeckBuilder:
    def __init__(self, spec: dict, base: Path):
        self.spec = spec
        self.base = base
        self.prs = Presentation()
        self.prs.slide_width = Inches(SLIDE_W)
        self.prs.slide_height = Inches(SLIDE_H)
        self.blank = self.prs.slide_layouts[6]
        self.qa_no = 0
        self.warnings: list[str] = []
        self.data = {}
        if spec.get("data"):
            import json
            dp = (base / spec["data"]).resolve()
            if dp.exists():
                self.data = json.loads(dp.read_text(encoding="utf-8"))
            else:
                self.warnings.append(f"데이터 파일 없음: {spec['data']} (fetch_data.py 먼저 실행)")

    # ---- 공통 요소
    def path(self, p) -> Path | None:
        if not p:
            return None
        q = (self.base / p).resolve()
        if not q.exists():
            self.warnings.append(f"이미지 없음: {p}")
            return None
        return q

    def new_slide(self, bg=BLACK):
        s = self.prs.slides.add_slide(self.blank)
        set_background(s, bg)
        return s

    def title_tab(self, slide, title: str):
        """왼쪽 위 흰색 탭 + 로고 + 제목."""
        logo = self.path(self.spec.get("logo"))
        if logo:
            lw, lh = fit(logo, LOGO_MAX_W, LOGO_H)
            text_x = LOGO_X + lw + 0.22
        else:
            logo_text = self.spec.get("logo_text") or self.spec.get("company", "")
            lw = text_width(logo_text, 18) + 0.1
            text_x = LOGO_X + lw + 0.2
        tw = text_width(title, 24) + 0.2
        tab_w = text_x + tw + 0.3
        tab = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(TAB_Y), Inches(tab_w), Inches(TAB_H))
        tab.name = "제목 탭"
        solid(tab, WHITE)
        no_line(tab)
        tab.shadow.inherit = False
        if logo:
            slide.shapes.add_picture(str(logo), Inches(LOGO_X), Inches(TAB_Y + (TAB_H - lh) / 2),
                                     Inches(lw), Inches(lh)).name = "로고"
        else:
            _, tf = text_box(slide, LOGO_X, TAB_Y, lw, TAB_H, anchor=MSO_ANCHOR.MIDDLE, wrap=False, margin=0)
            r = tf.paragraphs[0].add_run()
            r.text = logo_text
            set_run_font(r, FONT_TITLE, 18, self.spec.get("brand_color", "EA4335"))
        tb, tf = text_box(slide, text_x, TAB_Y, tw + 0.1, TAB_H, name="제목",
                          anchor=MSO_ANCHOR.MIDDLE, wrap=False, margin=0)
        r = tf.paragraphs[0].add_run()
        r.text = title
        set_run_font(r, FONT_TITLE, 24, BLACK)

    def body(self, slide, lines: list, x=BODY_X, y=BODY_Y, w=BODY_W, size=BODY_PT) -> float:
        """본문 글머리표 목록. 반환값: 본문이 끝나는 y(인치)."""
        if not lines:
            return y
        h_lines = 0
        for ln in lines:
            ln = str(ln)
            if ln.startswith("- "):
                h_lines += wrapped_lines(ln, size, w - 0.3)
            elif ln.strip() == "":
                h_lines += 0.75
            else:
                h_lines += wrapped_lines(ln, size, w - 0.3)
        height = h_lines * (size * 1.2 / 72 + 4 / 72) + 0.15
        tb, tf = text_box(slide, x, y, w, height, name="본문")
        first = True
        for ln in lines:
            ln = str(ln)
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            para_spacing(p)
            if ln.strip() == "":
                set_bullet(p, None, 0, 0)
                r = p.add_run()
                r.text = ""
                set_run_font(r, FONT_BODY, size * 0.6, WHITE)
                continue
            if ln.startswith("- "):
                set_bullet(p, None, 0.28, 0)
                add_rich_text(p, ln, size)
            else:
                set_bullet(p, "•", 0.28, -0.28)
                add_rich_text(p, ln, size)
        return y + height

    def caption(self, slide, text: str):
        tb, tf = text_box(slide, 6.0, 7.08, 7.1, 0.32, name="출처", margin=0)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = text
        set_run_font(r, FONT_BODY, 10, "8A8A8A")

    def notes(self, slide, text):
        """발표 대본. 애니메이션 표시 '(나)'·'(모핑)'은 넣지 않는다 (원고에 있어도 지운다)."""
        if not text:
            return
        lines = []
        for ln in str(text).strip().splitlines():
            ln = re.sub(r"\((나|모핑)\)", "", ln).rstrip()
            if ln.strip():
                lines.append(ln)
        slide.notes_slide.notes_text_frame.text = "\n".join(lines)

    # ---- 미디어 배치
    def place_media(self, slide, sd: dict, area: tuple[float, float, float, float]):
        x0, y0, w0, h0 = area
        items = []
        imgs = sd.get("images") or ([sd["image"]] if sd.get("image") else [])
        for p in imgs:
            q = self.path(p)
            if q:
                items.append(("image", q))
        charts = sd.get("charts") or ([sd["chart"]] if sd.get("chart") else [])
        for c in charts:
            items.append(("chart", c))
        if sd.get("table"):
            items.append(("table", sd["table"]))
        if not items or h0 < 0.6:
            return
        gap = 0.3
        n = len(items)
        if all(k == "image" for k, _ in items):
            # 같은 높이로 나란히, 전체 폭에 맞춘다.
            ratios = [image_size(q)[0] / image_size(q)[1] for _, q in items]
            h = min(h0, (w0 - gap * (n - 1)) / sum(ratios))
            total = h * sum(ratios) + gap * (n - 1)
            x = x0 + (w0 - total) / 2
            for (_, q), r in zip(items, ratios):
                slide.shapes.add_picture(str(q), Inches(x), Inches(y0 + (h0 - h) / 2), Inches(h * r), Inches(h))
                x += h * r + gap
            return
        cw = (w0 - gap * (n - 1)) / n
        for i, (kind, obj) in enumerate(items):
            cx = x0 + i * (cw + gap)
            if kind == "image":
                iw, ih = fit(obj, cw, h0)
                slide.shapes.add_picture(str(obj), Inches(cx + (cw - iw) / 2), Inches(y0 + (h0 - ih) / 2),
                                         Inches(iw), Inches(ih))
            elif kind == "chart":
                self.chart(slide, obj, (cx, y0, cw, h0))
            else:
                self.table(slide, obj, (cx, y0, cw, h0))

    # ---- 차트
    def resolve_chart(self, c: dict) -> dict:
        """source: price_1y | price_5y | annual.<지표> | quarterly.<지표> 이면 data.json 값으로 채운다.
        차트 항목에 직접 쓴 키가 자동 값보다 우선한다."""
        src = c.get("source")
        if not src:
            return c
        D = self.data
        if not D:
            raise ValueError(f"source '{src}'를 쓰려면 deck.yaml에 data: 경로가 필요합니다")
        if src in ("price_1y", "price_5y"):
            pts = D["price_daily_1y"] if src == "price_1y" else D["price_weekly_5y"]
            pr, v = D.get("price", {}), D.get("valuation", {})
            until = str(c.get("until") or "")
            if until:
                # 과거 기준일의 주가 화면 (F/U에서 '지난 분석 당시' 등). 시가총액 외 현재 기준 지표는 뺀다.
                import datetime as _dt
                pool = D.get("price_daily_2y") or D["price_weekly_5y"] if src == "price_1y" else D["price_weekly_5y"]
                days = 365 if src == "price_1y" else 365 * 5
                end = _dt.date.fromisoformat(until)
                start = (end - _dt.timedelta(days=days)).isoformat()
                pts = [p for p in pool if start <= p[0] <= until]
                if not pts:
                    raise ValueError(f"{until} 이전 주가 데이터가 없습니다 (fetch_data.py를 다시 실행)")
                hi, lo = max(p[1] for p in pts), min(p[1] for p in pts)
                v = {"market_cap": float(c["market_cap"])} if c.get("market_cap") else {}
                pr = {**pr, "high_52w": hi, "low_52w": lo}
            first, last = pts[0][1], pts[-1][1]
            # 축에는 매월(5년이면 매년) 첫 거래일만 이름을 붙인다. 나머지는 빈 이름.
            dates, prev = [], None
            for day, _ in pts:
                key = day[:7] if src == "price_1y" else day[:4]
                dates.append((day[2:7].replace("-", ".") if src == "price_1y" else day[:4]) if key != prev else "")
                prev = key
            stats = []
            if v.get("market_cap"):
                mc = v["market_cap"]
                stats.append(("시가총액", f"${mc / 1e12:,.2f}조" if mc >= 1e12 else f"${mc / 1e8:,.0f}억"))
            if v.get("trailing_pe"):
                stats.append(("P/E", f"{v['trailing_pe']:.1f}배"))
            if not until:
                stats.append(("배당수익률", f"{v['dividend_yield_pct']:.2f}%" if v.get("dividend_yield_pct") else "-"))
            if pr.get("high_52w"):
                stats.append(("52주 최고", f"{pr['high_52w']:,.2f}"))
                stats.append(("52주 최저", f"{pr['low_52w']:,.2f}"))
            if D.get("shares_outstanding") and not until:
                stats.append(("발행주식수", f"{D['shares_outstanding'] / 1e8:,.1f}억주"))
            auto = {
                "type": "line",
                "dates": dates,
                "values": [p[1] for p in pts],
                "name": f"{pr.get('long_name') or D.get('name')} ({D.get('ticker')}) · {pr.get('exchange', '')}",
                "price_text": f"{last:,.2f} {pr.get('currency', 'USD')}",
                "change_text": f"{last - first:+,.2f} ({(last / first - 1) * 100:+.1f}%) "
                               f"{'1년' if src == 'price_1y' else '5년'} · {pts[-1][0]} 종가",
                "stats": stats,
            }
        else:
            scope, metric = src.split(".", 1)
            series = D[scope][metric]
            until = str(c.get("until") or "9999")
            keys = sorted([k for k in series if series[k]["end"] <= until],
                          key=lambda k: series[k]["end"])[-int(c.get("n", 5 if scope == "annual" else 8)):]
            per_share = metric in ("eps_diluted", "eps_basic", "dps")
            shares = metric == "diluted_shares"
            scale = c.get("scale", 1 if per_share else 1e8)
            vals = [series[k]["value"] / scale for k in keys]
            if scope == "annual":
                cats = [str(k) for k in keys]
            else:
                cats = [f"{series[k]['q']}Q{series[k]['fy'] % 100:02d}" for k in keys]
            if per_share:
                labels = [("≈" if series[k].get("approx") else "") + f"${x:,.2f}" for k, x in zip(keys, vals)]
            elif shares:
                labels = [f"{x:,.1f}" for x in vals]
            else:
                labels = [f"${x:,.0f}억" if abs(x) >= 100 else (f"${x:,.1f}억" if abs(x) >= 10 else f"${x:,.2f}억")
                          for x in vals]
            auto = {"type": "bar", "categories": cats, "values": vals, "labels": labels}
        out = {**auto, **{k: v for k, v in c.items() if k != "source"}}
        # 전망치 덧붙이기 (예: 회사 CapEx 가이던스)
        if c.get("append_categories"):
            n0 = len(out["categories"])
            out["categories"] = list(out["categories"]) + [str(x) for x in c["append_categories"]]
            out["values"] = list(out["values"]) + [float(x) for x in c["append_values"]]
            out["labels"] = list(out.get("labels") or [None] * n0) + list(c.get("append_labels") or [None] * len(c["append_values"]))
            out["future"] = list(c.get("future") or []) + list(range(n0, len(out["categories"])))
        return out

    def chart(self, slide, c: dict, area):
        c = self.resolve_chart(c)
        kind = c.get("type", "bar")
        if kind == "line":
            return self.price_card(slide, c, area)
        return self.bar_chart(slide, c, area)

    def _chart_text(self, chart, size, color):
        chart.font.size = Pt(size)
        chart.font.color.rgb = rgb(color)
        chart.font.name = FONT_BODY
        for latin in chart._chartSpace.iter(qn("a:latin")):
            if latin.getparent().find(qn("a:ea")) is None:
                ea = etree.Element(qn("a:ea"))
                ea.set("typeface", FONT_BODY)
                latin.addnext(ea)

    def _no_fill_chart(self, chart):
        cs = chart._chartSpace
        for parent in (cs, cs.find(qn("c:chart")).find(qn("c:plotArea"))):
            spPr = parent.find(qn("c:spPr"))
            if spPr is None:
                spPr = etree.SubElement(parent, qn("c:spPr"))
                if parent is cs:
                    # c:spPr는 c:chart 다음, c:txPr 앞에 와야 한다.
                    cs.find(qn("c:chart")).addnext(spPr)
            for ch in list(spPr):
                spPr.remove(ch)
            etree.SubElement(spPr, qn("a:noFill"))
            ln = etree.SubElement(spPr, qn("a:ln"))
            etree.SubElement(ln, qn("a:noFill"))

    def bar_chart(self, slide, c: dict, area):
        """CapEx 슬라이드 같은 하늘색 막대그래프 (검정 배경 위, 값 축 없음)."""
        x, y, w, h = area
        cats = [str(v) for v in c["categories"]]
        vals = [float(v) for v in c["values"]]
        cd = CategoryChartData()
        cd.categories = cats
        cd.add_series(c.get("name", "값"), vals)
        title = c.get("title")
        gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y),
                                    Inches(w), Inches(h), cd)
        gf.name = c.get("title") or "막대그래프"
        ch = gf.chart
        self._no_fill_chart(ch)
        highlights = c.get("highlight") or []
        futures = c.get("future") or []
        labels = c.get("labels") or [None] * len(vals)
        tops = c.get("top_labels") or [None] * len(vals)
        fmt = c.get("format", "{:,.0f}")
        texts = [str(labels[i]) if labels[i] is not None else fmt.format(v) for i, v in enumerate(vals)]
        base = c.get("font_size", (18 if w > 7 else 14) if len(cats) <= 6 else 12)
        # 숫자 라벨·항목 이름이 막대 한 칸(slot) 폭 안에 들어가도록 각각 글꼴을 줄인다 (차트 제목은 그대로)
        slot = w * 0.9 / max(len(vals), 1)

        def fitted(strings):
            widest = max([text_width(str(t), base) for t in strings if t] + [0.01])
            return base if widest <= slot * 0.95 else max(8, int(base * slot * 0.95 / widest))

        size = fitted(texts + [t for t in tops if t])
        cat_size = fitted(cats)
        self._chart_text(ch, size, SKY)
        ch.has_legend = False
        if title:
            ch.has_title = True
            tf = ch.chart_title.text_frame
            tf.text = ""
            r = tf.paragraphs[0].add_run()
            r.text = title
            set_run_font(r, FONT_BODY, base, WHITE)
        else:
            ch.has_title = False
            ac = ch._chartSpace.find(qn("c:chart")).find(qn("c:autoTitleDeleted"))
            if ac is not None:
                ac.set("val", "1")
        plot = ch.plots[0]
        plot.gap_width = c.get("gap", 110)
        plot.vary_by_categories = False
        va = ch.value_axis
        va.visible = False
        va.has_major_gridlines = False
        va.minimum_scale = 0
        # 가장 높은 막대 위에도 라벨이 들어갈 자리를 만든다 (모자라면 PowerPoint가 라벨을 막대 안으로 넣어 가려짐)
        line_in = size * 1.25 / 72
        plot_h = h - (base * 1.8 / 72 if title else 0) - cat_size * 1.6 / 72 - 0.15
        need = 0.0
        for i, v in enumerate(vals):
            lab_h = line_in * (2 if tops[i] else 1) + 0.08
            if v > 0 and plot_h > lab_h:
                need = max(need, v * plot_h / (plot_h - lab_h))
        top = max(need * 1.04, float(c.get("max") or 0))
        if top > 0:
            va.maximum_scale = top
        ca = ch.category_axis
        ca.format.line.fill.background()
        ca.tick_labels.font.size = Pt(cat_size)
        ca.tick_labels.font.color.rgb = rgb(SKY)
        ca.has_major_gridlines = False
        ser = plot.series[0]
        solid(ser.format, SKY)
        ser.format.line.fill.background()
        for i, v in enumerate(vals):
            pt = ser.points[i]
            color = SKY
            if i in futures or cats[i] in [str(f) for f in futures]:
                color = PEACH
            if i in highlights or cats[i] in [str(f) for f in highlights]:
                color = YELLOW
            solid(pt.format, color)
            pt.format.line.fill.background()
            dl = pt.data_label
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
            tf = dl.text_frame
            tf.text = ""
            p = tf.paragraphs[0]
            if tops[i]:
                r = p.add_run()
                r.text = str(tops[i])
                set_run_font(r, FONT_BODY, size, YELLOW)
                p = tf.add_paragraph()
            lab = texts[i]
            if lab != "":
                r = p.add_run()
                r.text = str(lab)
                set_run_font(r, FONT_BODY, size, color)
            body_pr = tf._txBody.find(qn("a:bodyPr"))
            if body_pr is not None:
                body_pr.set("wrap", "none")   # 라벨이 막대 폭에서 줄바꿈되지 않게

    def price_card(self, slide, c: dict, area):
        """구글 금융 화면처럼 흰 카드 안에 1년 주가 선그래프 + 주요 지표."""
        x, y, w, h = area
        card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        card.name = "주가 카드"
        solid(card, WHITE)
        no_line(card)
        card.shadow.inherit = False
        pad = 0.25
        # 머리글
        _, tf = text_box(slide, x + pad, y + 0.15, w - 2 * pad, 1.15, margin=0)
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = c.get("name", "")
        set_run_font(r, FONT_BODY, 14, DARK)
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = c.get("price_text", "")
        set_run_font(r, FONT_TITLE, 26, DARK)
        chg = c.get("change_text", "")
        up = not chg.strip().startswith("-")
        line_color = UP_RED if up else DOWN_BLUE
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = chg
        set_run_font(r, FONT_BODY, 13, line_color)
        stats = c.get("stats") or []
        stats_h = 0.32 * ((len(stats) + 2) // 3) + 0.15 if stats else 0
        # 선그래프
        cd = CategoryChartData()
        cd.categories = [str(d) for d in c["dates"]]
        cd.add_series("종가", [float(v) for v in c["values"]])
        cy = y + 1.35
        chh = h - 1.35 - stats_h - 0.1
        gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(x + 0.1), Inches(cy), Inches(w - 0.2),
                                    Inches(chh), cd)
        gf.name = "주가 그래프"
        ch = gf.chart
        self._no_fill_chart(ch)
        self._chart_text(ch, 10, "70757A")
        ch.has_legend = False
        ch.has_title = False
        atd = ch._chartSpace.find(qn("c:chart")).find(qn("c:autoTitleDeleted"))
        if atd is not None:
            atd.set("val", "1")
        ser = ch.plots[0].series[0]
        ser.smooth = False
        ser.format.line.color.rgb = rgb(line_color)
        ser.format.line.width = Pt(1.75)
        ser.marker.style = XL_MARKER_STYLE.NONE
        va = ch.value_axis
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = rgb("E8EAED")
        va.format.line.fill.background()
        vals = [float(v) for v in c["values"]]
        lo, hi = min(vals), max(vals)
        span = hi - lo or 1
        va.minimum_scale = max(0, round(lo - span * 0.1, -1 if lo > 50 else 0))
        va.tick_labels.font.size = Pt(10)
        ca = ch.category_axis
        ca.format.line.color.rgb = rgb("DADCE0")
        ca.tick_labels.font.size = Pt(10)
        n = len(c["dates"])
        cax = ca._element
        labelled = sum(1 for x in c["dates"] if x)
        skip = 1 if labelled < n / 3 else max(1, n // 6)   # 빈 이름 방식이면 모두 표시
        for tag, val in (("c:tickLblSkip", skip), ("c:tickMarkSkip", max(skip, n // 12 or 1))):
            el = cax.find(qn(tag))
            if el is None:
                el = etree.SubElement(cax, qn(tag))
            el.set("val", str(val))
        # tickLblSkip/tickMarkSkip/noMultiLvlLbl 순서를 스키마에 맞춘다.
        nml = cax.find(qn("c:noMultiLvlLbl"))
        if nml is not None:
            cax.remove(nml)
            cax.append(nml)
        # 지표 표
        if stats:
            sy = y + h - stats_h
            colw = (w - 2 * pad) / 3
            for i, (k, v) in enumerate(stats):
                col, row = i % 3, i // 3
                _, tf = text_box(slide, x + pad + col * colw, sy + row * 0.32, colw - 0.1, 0.3, margin=0)
                p = tf.paragraphs[0]
                r = p.add_run()
                r.text = f"{k}  "
                set_run_font(r, FONT_BODY, 11, "70757A")
                r = p.add_run()
                r.text = str(v)
                set_run_font(r, FONT_BODY, 11, DARK)

    # ---- 표
    def table(self, slide, t: dict, area):
        x, y, w, h = area
        style = t.get("style", "grid")
        header = t.get("header")
        rows = t.get("rows", [])
        data = ([header] if header else []) + rows
        nr, nc = len(data), max(len(r) for r in data)
        size = t.get("font_size", 16 if style == "grid" else 11)
        row_h = t.get("row_height", size * 1.9 / 72)
        th = min(h, row_h * nr)
        tw = t.get("width", w)
        tx = x + (w - tw) / 2
        if style == "statement":
            card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(tx), Inches(y), Inches(tw), Inches(th + 0.3))
            solid(card, WHITE)
            no_line(card)
            card.shadow.inherit = False
            tx, y, tw = tx + 0.15, y + 0.15, tw - 0.3
        gf = slide.shapes.add_table(nr, nc, Inches(tx), Inches(y), Inches(tw), Inches(th))
        gf.name = t.get("name", "표")
        tbl = gf.table
        tblPr = tbl._tbl.tblPr
        for attr in ("firstRow", "bandRow"):
            tblPr.set(attr, "0")
        sid = tblPr.find(qn("a:tableStyleId"))
        if sid is not None:
            sid.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"   # 스타일 없음, 표 눈금
        widths = t.get("col_widths")
        if widths:
            tot = sum(widths)
            for i, cw in enumerate(widths):
                tbl.columns[i].width = Inches(tw * cw / tot)
        else:
            first = (0.5 if nc <= 3 else 0.36) if style == "statement" else 1 / nc
            rest = (1 - first) / max(nc - 1, 1)
            for i in range(nc):
                tbl.columns[i].width = Inches(tw * (first if i == 0 else rest))
        for i in range(nr):
            tbl.rows[i].height = Inches(row_h)
        for r_i, row in enumerate(data):
            is_head = header is not None and r_i == 0
            for c_i in range(nc):
                cell = tbl.cell(r_i, c_i)
                val = str(row[c_i]) if c_i < len(row) and row[c_i] is not None else ""
                cell.margin_left = cell.margin_right = Inches(0.06)
                cell.margin_top = cell.margin_bottom = Inches(0.02)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                fill = WHITE
                if style == "grid" and (is_head or (t.get("head_col") and c_i == 0)):
                    fill = TABLE_HEAD
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(fill)
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                if style == "grid":
                    p.alignment = PP_ALIGN.CENTER
                else:
                    p.alignment = PP_ALIGN.LEFT if c_i == 0 else PP_ALIGN.RIGHT
                face = FONT_TITLE if (is_head or style == "grid") else FONT_BODY
                add_rich_text(p, val, size, DARK if style == "statement" else BLACK, face=face, on_light=True)
                self._cell_borders(cell, style, is_head, r_i == nr - 1)
        for m in t.get("merge", []):
            r1, c1, r2, c2 = m
            tbl.cell(r1, c1).merge(tbl.cell(r2, c2))
        # 노란 화살표로 주목할 행 표시 (예시의 연도별 실적·주요 주주 슬라이드)
        for r_i in t.get("arrows", []):
            ay = y + row_h * r_i + row_h / 2 - 0.13
            ar = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(tx - 0.65 - (0.15 if style == "statement" else 0)),
                                        Inches(ay), Inches(0.45), Inches(0.26))
            solid(ar, YELLOW)
            no_line(ar)
            ar.shadow.inherit = False

    @staticmethod
    def _cell_borders(cell, style, is_head, is_last):
        tcPr = cell._tc.get_or_add_tcPr()
        for side in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
            old = tcPr.find(qn(side))
            if old is not None:
                tcPr.remove(old)
        specs = {}
        if style == "grid":
            specs = {s: ("000000", 12700) for s in ("a:lnL", "a:lnR", "a:lnT", "a:lnB")}
        else:
            specs = {"a:lnL": None, "a:lnR": None, "a:lnT": None,
                     "a:lnB": ("000000", 9525) if is_head else ("E0E0E0", 6350)}
        fill = tcPr.find(qn("a:solidFill"))
        for side in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
            ln = etree.Element(qn(side))
            spec = specs.get(side)
            if spec:
                ln.set("w", str(spec[1]))
                sf = etree.SubElement(ln, qn("a:solidFill"))
                c = etree.SubElement(sf, qn("a:srgbClr"))
                c.set("val", spec[0])
            else:
                ln.set("w", "0")
                etree.SubElement(ln, qn("a:noFill"))
            if fill is not None:
                fill.addprevious(ln)
            else:
                tcPr.append(ln)

    # ---- 슬라이드 종류
    def s_cover(self, sd):
        s = self.new_slide()
        img = self.path(sd.get("image") or self.spec.get("cover_image"))
        name = sd.get("name") or self.spec.get("company", "")
        sub = sd.get("subtitle", "기업분석")
        if img:
            iw, ih = fit(img, 11.6, 3.7)
            s.shapes.add_picture(str(img), Inches((SLIDE_W - iw) / 2), Inches(0.35 + (3.7 - ih) / 2),
                                 Inches(iw), Inches(ih))
            ty = 4.2
        else:
            ty = 2.4
        big, small = 140, 84
        total = text_width(name, big) + 0.6 + text_width(sub, small)
        if total > 11.6:
            k = 11.6 / total
            big, small = big * k, small * k
        _, tf = text_box(s, 0.9, ty, 11.8, 2.9, name="제목", anchor=MSO_ANCHOR.BOTTOM, wrap=False)
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = name
        set_run_font(r, FONT_DISPLAY, big, WHITE)
        r = p.add_run()
        r.text = "   " + sub
        set_run_font(r, FONT_DISPLAY, small, WHITE)
        self.notes(s, sd.get("notes"))
        return s

    def s_section(self, sd):
        s = self.new_slide(WHITE)
        _, tf = text_box(s, 0.9, 2.2, 11.5, 2.4, name="제목", anchor=MSO_ANCHOR.MIDDLE)
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = sd.get("text") or sd.get("title", "")
        set_run_font(r, FONT_DISPLAY, sd.get("size", 88), BLACK)
        if sd.get("subtitle"):
            p = tf.add_paragraph()
            r = p.add_run()
            r.text = sd["subtitle"]
            set_run_font(r, FONT_BODY, 24, "5F6368")
        self.notes(s, sd.get("notes"))
        return s

    def s_content(self, sd):
        s = self.new_slide()
        self.title_tab(s, sd.get("title", ""))
        layout = sd.get("layout", "below")
        lines = sd.get("lines") or []
        if layout == "right":
            tw = sd.get("text_width", 6.2)
            self.body(s, lines, w=tw)
            mx = BODY_X + tw + 0.3
            self.place_media(s, sd, (mx, 1.0, SLIDE_W - mx - 0.5, CONTENT_BOTTOM - 1.0))
        elif layout == "left":
            mw = sd.get("media_width", 5.6)
            self.place_media(s, sd, (0.6, 1.1, mw, CONTENT_BOTTOM - 1.1))
            tx = 0.6 + mw + 0.4
            self.body(s, lines, x=tx, y=1.3, w=SLIDE_W - tx - 0.4)
        else:
            end = self.body(s, lines)
            top = end + 0.2 if lines else 1.0
            self.place_media(s, sd, (BODY_X, top, BODY_W, CONTENT_BOTTOM - top))
        if sd.get("source"):
            self.caption(s, "출처: " + sd["source"])
        self.notes(s, sd.get("notes"))
        return s

    def s_qa(self, sd):
        self.qa_no += 1
        n = sd.get("n", self.qa_no)
        s = self.new_slide()
        self.title_tab(s, sd.get("title", "Q & A"))
        q_lines = wrapped_lines(f"Q{n}. " + sd["q"], BODY_PT, BODY_W - 0.6)
        answers = sd["a"] if isinstance(sd["a"], list) else [sd["a"]]
        a_lines = sum(wrapped_lines(a, BODY_PT, BODY_W - 0.6) for a in answers)
        h = (q_lines + a_lines + 1) * (BODY_PT * 1.25 / 72) + 0.3
        _, tf = text_box(s, BODY_X, BODY_Y, BODY_W, h, name="본문")
        p = tf.paragraphs[0]
        set_bullet(p, None, 0.6, -0.6)
        r = p.add_run()
        r.text = f"Q{n}.  "
        set_run_font(r, FONT_BODY, BODY_PT, WHITE)
        add_rich_text(p, sd["q"], BODY_PT)
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = ""
        set_run_font(r, FONT_BODY, 12, WHITE)
        for i, a in enumerate(answers):
            p = tf.add_paragraph()
            para_spacing(p, 2, 1.05)
            set_bullet(p, None, 0.6, -0.6 if i == 0 else 0)
            if i == 0:
                r = p.add_run()
                r.text = "A.   "
                set_run_font(r, FONT_BODY, BODY_PT, WHITE)
            add_rich_text(p, a, BODY_PT)
        if sd.get("source"):
            self.caption(s, "출처: " + sd["source"])
        self.notes(s, sd.get("notes"))
        return s

    def s_segments(self, sd):
        """매출 구성 구조도: 왼쪽 괄호선 → 사업부 상자 → 세부 매출원 상자."""
        s = self.new_slide()
        self.title_tab(s, sd.get("title", "매출 구성"))
        segs = sd["segments"]
        n = len(segs)
        top, bottom = 1.3, 6.2 if not sd.get("images") else 5.9
        slot = (bottom - top) / n
        box_w, box_h = sd.get("box_width", 3.6), min(1.05, slot * 0.72)
        bx = 2.6
        root = (0.95, top + (bottom - top) / 2)
        for i, seg in enumerate(segs):
            cy = top + slot * i + slot / 2
            ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(root[0]), Inches(root[1]),
                                        Inches(bx - 0.25), Inches(cy))
            ln.line.color.rgb = rgb(YELLOW)
            ln.line.width = Pt(3)
            box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(bx), Inches(cy - box_h / 2),
                                     Inches(box_w), Inches(box_h))
            box.name = f"사업부 {i + 1}"
            solid(box, LIME)
            no_line(box)
            box.shadow.inherit = False
            tf = box.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            add_rich_text(p, seg["name"], sd.get("font_size", 24), BLACK, face=FONT_TITLE, on_light=True)
            items = seg.get("items") or []
            if items:
                ix = bx + box_w + 1.2
                iw = SLIDE_W - ix - 0.5
                ih = min(0.62, (slot * 0.95) / len(items) - 0.08)
                span = len(items) * (ih + 0.1) - 0.1
                for j, it in enumerate(items):
                    iy = cy - span / 2 + j * (ih + 0.1)
                    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(bx + box_w + 0.2), Inches(cy),
                                                Inches(ix - 0.15), Inches(iy + ih / 2))
                    ln.line.color.rgb = rgb(YELLOW)
                    ln.line.width = Pt(3)
                    ib = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(ix), Inches(iy), Inches(iw), Inches(ih))
                    solid(ib, LIME)
                    no_line(ib)
                    ib.shadow.inherit = False
                    tf = ib.text_frame
                    tf.word_wrap = True
                    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                    p = tf.paragraphs[0]
                    p.alignment = PP_ALIGN.CENTER
                    add_rich_text(p, it, 18, BLACK, face=FONT_TITLE, on_light=True)
        if sd.get("images"):
            self.place_media(s, {"images": sd["images"]}, (6.9, 5.95, 5.9, 1.2))
        if sd.get("source"):
            self.caption(s, "출처: " + sd["source"])
        self.notes(s, sd.get("notes"))
        return s

    def s_end(self, sd):
        s = self.new_slide()
        img = self.path(sd.get("image") or self.spec.get("cover_image"))
        if img:
            iw, ih = fit(img, 9.5, 3.6)
            s.shapes.add_picture(str(img), Inches((SLIDE_W - iw) / 2), Inches(0.9 + (3.6 - ih) / 2),
                                 Inches(iw), Inches(ih))
        _, tf = text_box(s, 0.5, 4.9, SLIDE_W - 1.0, 1.6, name="제목", anchor=MSO_ANCHOR.MIDDLE)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = sd.get("text", "감사합니다")
        set_run_font(r, FONT_DISPLAY, 66, WHITE)
        self.notes(s, sd.get("notes"))
        return s

    # ================================================================ F/U(실적 리뷰) 양식
    def stamp(self, slide, date: str):
        """오른쪽 위 기준일 표기: (2026-05-10)."""
        _, tf = text_box(slide, 9.9, TAB_Y + 0.05, 3.0, 0.38, name="기준일", anchor=MSO_ANCHOR.MIDDLE, margin=0)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = date if str(date).startswith("(") else f"({date})"
        set_run_font(r, FONT_LIGHT, 16, WHITE)

    def s_fu_cover(self, sd):
        """F/U 표지: 주제에 맞는 배경 사진 전면 + 회사명(노랑) + 한 줄 후킹 문구 + 오른쪽 아래 작은 로고."""
        s = self.new_slide(sd.get("bg", "342E7F"))
        img = self.path(sd.get("image") or self.spec.get("fu_cover_image"))
        if img:
            add_fill_picture(s, img, 0, 0, SLIDE_W, SLIDE_H, name="배경")
            dim = float(sd.get("dim", 0.35))     # 밝은 사진이면 검은 막을 덮어 글자를 살린다 (0=없음)
            if dim > 0:
                veil = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SLIDE_W), Inches(SLIDE_H))
                veil.name = "어둡게"
                solid(veil, BLACK)
                no_line(veil)
                veil.shadow.inherit = False
                clr = veil.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
                alpha = etree.SubElement(clr, qn("a:alpha"))
                alpha.set("val", str(int(dim * 100000)))
        name = sd.get("name") or self.spec.get("company", "")
        hook = sd.get("hook", "")
        name_pt = sd.get("name_size", 170)
        name_pt = min(name_pt, name_pt * 11.4 / max(text_width(name, name_pt), 0.1))
        hook_pt = sd.get("hook_size", 140)
        plain = MARK_RE.sub(lambda m: m.group(0)[2:-2], hook)
        hook_pt = min(hook_pt, hook_pt * 11.4 / max(text_width(plain, hook_pt) * 0.85, 0.1))
        _, tf = text_box(s, 1.0, 0.45, 11.6, 6.5, name="제목", anchor=MSO_ANCHOR.MIDDLE)
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = name
        set_run_font(r, FONT_DISPLAY, name_pt, YELLOW)
        add_glow(r)
        if hook:
            p = tf.add_paragraph()
            p.space_before = Pt(name_pt * 0.35)
            for tok in MARK_RE.split(hook):
                if not tok:
                    continue
                big = tok.startswith("[[")
                r = p.add_run()
                r.text = tok[2:-2] if big else tok
                set_run_font(r, FONT_DISPLAY, hook_pt if big else hook_pt * 0.7, YELLOW if big else WHITE)
                add_glow(r)
        icon = self.path(sd.get("logo_small") or self.spec.get("logo_small"))
        if icon:
            iw, ih = fit(icon, 0.9, 0.6)
            s.shapes.add_picture(str(icon), Inches(SLIDE_W - 0.25 - iw), Inches(SLIDE_H - 0.25 - ih),
                                 Inches(iw), Inches(ih)).name = "작은 로고"
        self.notes(s, sd.get("notes"))
        return s

    def s_full_image(self, sd):
        """사진 한 장을 화면 가득 (실적발표 자료 표지, 행사 사진, 제품 사진 등). 제목 탭은 선택."""
        s = self.new_slide(sd.get("bg", BLACK))
        img = self.path(sd.get("image"))
        if img:
            if sd.get("fit"):      # 자르지 않고 가운데 맞춤 (표·도표 이미지)
                iw, ih = fit(img, SLIDE_W, SLIDE_H)
                s.shapes.add_picture(str(img), Inches((SLIDE_W - iw) / 2), Inches((SLIDE_H - ih) / 2),
                                     Inches(iw), Inches(ih))
            else:
                add_fill_picture(s, img, 0, 0, SLIDE_W, SLIDE_H, name="전면 사진")
        if sd.get("tab"):          # 마지막 장처럼 흰 탭에 글자만 (로고 없음)
            tw = text_width(sd["tab"], 24) + 0.6
            tab = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(TAB_Y), Inches(tw), Inches(TAB_H))
            solid(tab, WHITE)
            no_line(tab)
            tab.shadow.inherit = False
            tf = tab.text_frame
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = sd["tab"]
            set_run_font(r, FONT_TITLE, 24, BLACK)
        elif sd.get("title"):
            self.title_tab(s, sd["title"])
        if sd.get("source"):
            self.caption(s, "출처: " + sd["source"])
        self.notes(s, sd.get("notes"))
        return s

    def s_quote(self, sd):
        """CEO 한마디: 오른쪽에 인물 사진(위아래 꽉), 왼쪽에 큰 따옴표 인용문. 핵심어 [[ ]]는 노랑·더 크게."""
        s = self.new_slide()
        img = self.path(sd.get("image"))
        pw = sd.get("photo_width", 6.5)
        if img:
            add_fill_picture(s, img, SLIDE_W - pw, 0, pw, SLIDE_H, name="인물 사진")
        lines = sd["quote"] if isinstance(sd["quote"], list) else [sd["quote"]]
        tw = sd.get("text_width", 8.2)
        size = sd.get("size", 40)
        widest = max(text_width(MARK_RE.sub(lambda m: m.group(0)[2:-2], ln), size) * 1.08 + 0.6 for ln in lines)
        if widest > tw:
            size = max(26, int(size * tw / widest))
        _, tf = text_box(s, 1.1, 0.9, tw, 5.4, name="인용문", anchor=MSO_ANCHOR.MIDDLE)
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.line_spacing = 1.05
            text = ("“ " if i == 0 else "   ") + ln + (" ”" if i == len(lines) - 1 else "")
            for tok in MARK_RE.split(text):
                if not tok:
                    continue
                key = tok.startswith("[[")
                r = p.add_run()
                r.text = tok[2:-2] if key else tok
                set_run_font(r, FONT_BODY, size * 1.2 if key else size, YELLOW if key else WHITE)
                add_glow(r, rad_pt=6)
        if sd.get("by"):
            p = tf.add_paragraph()
            p.space_before = Pt(18)
            r = p.add_run()
            r.text = "— " + sd["by"]
            set_run_font(r, FONT_LIGHT, 18, "BFBFBF")
            add_glow(r, rad_pt=6)
        self.notes(s, sd.get("notes"))
        return s

    def s_guidance(self, sd):
        """다음 분기 가이던스 차트: 최근 분기 실적 + 가이던스(살구색) 막대, 막대 사이 성장률(노랑),
        막대 안 빨간 점 = 영업이익률, 초록 점 = EPS. 예시 덱 '2026-4분기 가이던스' 슬라이드 재현."""
        s = self.new_slide()
        self.title_tab(s, sd.get("title", "가이던스"))
        q = [str(x) for x in sd["quarters"]]
        vals = [float(v) for v in sd["revenue"]]
        n = len(q)
        rev_lab = sd.get("revenue_labels") or [f"{v:,.0f}" for v in vals]
        growth = sd.get("growth") or []
        margin = sd.get("margin") or []
        eps = sd.get("eps") or []
        guide = set(sd.get("guidance_index", [n - 1]))
        leg = {"bar": "매출", "margin": "Non-GAAP 영업이익률", "eps": "Non-GAAP 희석 EPS", **(sd.get("legend") or {})}
        # 범례 (위쪽 한 줄)
        lx, ly = 1.3, 1.05
        for col in (SKY, PEACH):
            b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(lx), Inches(ly + 0.1), Inches(0.5), Inches(0.17))
            solid(b, col); no_line(b); b.shadow.inherit = False
            lx += 0.58
        items = [(None, leg["bar"], SKY)]
        if margin:
            items.append((METRIC_RED, leg["margin"], METRIC_RED))
        if eps:
            items.append((METRIC_GREEN, leg["eps"], METRIC_GREEN))
        for dot, text, col in items:
            if dot:
                d = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(lx), Inches(ly + 0.1), Inches(0.18), Inches(0.18))
                solid(d, dot); no_line(d); d.shadow.inherit = False
                lx += 0.25
            _, tf = text_box(s, lx, ly, text_width(text, 18) + 0.2, 0.38, wrap=False, margin=0)
            r = tf.paragraphs[0].add_run(); r.text = text
            set_run_font(r, FONT_BODY, 18, col)
            lx += text_width(text, 18) + 0.55
        # 막대
        base_y, max_h, min_h = 6.0, 3.4, 1.75   # 가장 높은 막대 위 성장률·값 라벨이 범례와 겹치지 않게
        x0, span = 1.2, 10.9
        slot = span / n
        bw = min(0.9, slot * 0.45)
        vmax = max(vals)
        lab_pt = 18
        widest = max(text_width(t, lab_pt) for t in rev_lab + growth + q)
        if widest > slot * 0.95:
            lab_pt = max(11, int(lab_pt * slot * 0.95 / widest))
        tops = []
        for i, v in enumerate(vals):
            h = max(min_h, max_h * v / vmax)
            cx = x0 + slot * (i + 0.5)
            top = base_y - h
            tops.append(top)
            col = PEACH if i in guide else SKY
            b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(cx - bw / 2), Inches(top), Inches(bw), Inches(h))
            b.name = f"막대 {q[i]}"
            solid(b, col); no_line(b); b.shadow.inherit = False
            for txt, yy, colr, pt in ((rev_lab[i], top - 0.42, col, lab_pt), (q[i], base_y + 0.04, col, lab_pt)):
                _, tf = text_box(s, cx - slot / 2, yy, slot, 0.38, wrap=False, margin=0)
                p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
                r = p.add_run(); r.text = txt
                set_run_font(r, FONT_BODY, pt, colr)
            # 막대 안 지표 점: 빨강(영업이익률) 위, 초록(EPS) 아래
            dot_pt = min(16, lab_pt)
            if i < len(margin) and margin[i]:
                yy = top + 0.18
                _, tf = text_box(s, cx - slot / 2, yy, slot, 0.32, wrap=False, margin=0)
                p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
                r = p.add_run(); r.text = str(margin[i]); set_run_font(r, FONT_BODY, dot_pt, METRIC_RED)
                d = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - 0.105), Inches(yy + 0.36), Inches(0.21), Inches(0.21))
                solid(d, METRIC_RED); no_line(d); d.shadow.inherit = False
            if i < len(eps) and eps[i]:
                yy = top + 0.18 + (0.66 if margin else 0)
                d = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - 0.095), Inches(yy), Inches(0.19), Inches(0.17))
                solid(d, METRIC_GREEN); no_line(d); d.shadow.inherit = False
                lines = str(eps[i]).split("\n")
                _, tf = text_box(s, cx - slot / 2, yy + 0.2, slot, 0.3 * len(lines) + 0.05, margin=0)
                for k, ln in enumerate(lines):
                    p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                    p.alignment = PP_ALIGN.CENTER
                    r = p.add_run(); r.text = ln; set_run_font(r, FONT_BODY, dot_pt, METRIC_GREEN)
        # 막대 사이 성장률 (노랑): 다음 막대 숫자 높이, 두 막대 사이 가운데
        for i, g in enumerate(growth[: n - 1]):
            if not g:
                continue
            cx = x0 + slot * (i + 1)
            yy = min(tops[i], tops[i + 1]) - 0.88
            _, tf = text_box(s, cx - slot / 2, yy, slot, 0.38, wrap=False, margin=0)
            p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
            r = p.add_run(); r.text = str(g); set_run_font(r, FONT_BODY, lab_pt, YELLOW)
        if sd.get("source"):
            self.caption(s, "출처: " + sd["source"])
        self.notes(s, sd.get("notes"))
        return s

    def build(self, out: Path):
        kinds = {"cover": self.s_cover, "section": self.s_section, "qa": self.s_qa,
                 "segments": self.s_segments, "end": self.s_end,
                 "fu_cover": self.s_fu_cover, "full_image": self.s_full_image,
                 "quote": self.s_quote, "guidance": self.s_guidance}
        for i, sd in enumerate(self.spec["slides"], 1):
            fn = kinds.get(sd.get("type", "content"), self.s_content)
            try:
                slide = fn(sd)
                if sd.get("date"):
                    self.stamp(slide, sd["date"])
            except Exception as e:  # 어느 슬라이드에서 문제가 났는지 알려준다
                raise SystemExit(f"슬라이드 {i} ({sd.get('title') or sd.get('type')}) 생성 실패: {e}")
        cp = self.prs.core_properties
        cp.title = self.spec.get("deck_title") or f"{self.spec.get('company', '')} 기업분석"
        cp.author = self.spec.get("author", "")
        out.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(out)
        return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("spec", type=Path, help="deck.yaml 경로")
    ap.add_argument("-o", "--out", type=Path, help="출력 .pptx (기본: output/<ticker>_기업분석.pptx)")
    a = ap.parse_args()
    spec = yaml.safe_load(a.spec.read_text(encoding="utf-8"))
    here = Path(__file__).resolve().parent
    out = a.out or here / "output" / spec.get("output", f"{spec.get('ticker', 'deck')}_기업분석.pptx")
    b = DeckBuilder(spec, a.spec.resolve().parent)
    b.build(out)
    for w in b.warnings:
        print("경고:", w, file=sys.stderr)
    print(f"저장: {out}  ({len(spec['slides'])}장)")


if __name__ == "__main__":
    main()
